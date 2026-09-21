"""Livelli di riservatezza (§134), export e stampa riservata (§135), azioni amministrative sensibili (§136)."""
import io
import json
import time
import zipfile
from datetime import datetime, timedelta

import pytest
import pytest_asyncio
from docx import Document as WordDocument
from pypdf import PdfReader, PdfWriter
from sqlalchemy import select

from backend.auth.privacypa_confidentiality import step_up_satisfied, validate_config
from backend.auth.portal_jwt import TokenPayload
from backend.models.privacypa_audit import AuditLog
from backend.models.privacypa_identity import PrivacyRoleCatalog, PrivacyRoleCode as R, ScopeType
from backend.tests.conftest import TA, TB, _h, _identity, ifm

API = "/api/privacypa"
DOCS, BR, PR, DPIAS, POLICY = f"{API}/documents", f"{API}/breaches", f"{API}/projects", f"{API}/dpias", f"{API}/policy/confidentiality"
ADMIN = _h(TA, "root", "admin")
HR, RE = "HIGHLY_RESTRICTED", "RESTRICTED"


@pytest_asyncio.fixture
async def w(env):
    grant = env["grant"]
    for user, role, scope, unit in (
        ("mgr", R.PRIVACY_MANAGER, ScopeType.TENANT, None), ("dpo", R.DPO, ScopeType.TENANT, None),
        ("op", R.OPERATOR, ScopeType.TENANT, None), ("reader", R.READ_ONLY, ScopeType.TENANT, None),
        ("aud", R.AUDITOR, ScopeType.TENANT, None), ("ict", R.ICT, ScopeType.TENANT, None),
        ("ref-a", R.PRIVACY_REFERENT, ScopeType.UNIT, "unit-a"), ("ref-b", R.PRIVACY_REFERENT, ScopeType.UNIT, "unit-b"),
    ):
        await grant(TA, user, role, scope, unit)
    await grant(TB, "b-mgr", R.PRIVACY_MANAGER)
    return {k: _h(TA, k) for k in ("mgr", "dpo", "op", "reader", "aud", "ict", "ref-a", "ref-b")} | {"b": _h(TB, "b-mgr")}


def strong(sub: str, *, tenant: str = TA, tenant_role: str | None = None, age: float = 0, amr=("pwd", "otp"), key: str = "auth_time") -> dict:
    """Identità con una autenticazione a più fattori di `age` secondi fa (claim come li emetterebbe il Portal/IdP)."""
    ident = json.loads(_identity(tenant, sub, tenant_role))
    ident["amr"], ident[key] = list(amr), time.time() - age
    return {"X-Test-User": json.dumps(ident)}


async def _set_policy(env, config: dict):
    c = env["client"]
    version = (await c.get(POLICY, headers=ADMIN)).json()["row_version"]
    r = await c.put(POLICY, json=config, headers=ifm(version, ADMIN))
    assert r.status_code == 200, r.text
    return r.json()


def _pdf() -> bytes:
    out = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(595, 842)
    writer.write(out)
    return out.getvalue()


def _docx() -> bytes:
    doc = WordDocument()
    doc.add_paragraph("Contenuto riservato")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def _document(env, who, level=None, *, title="Relazione", link=True):
    body = {"title": title, "document_type": "OTHER"}
    if level:
        body["confidentiality_level"] = level
    if link:
        body["links"] = [{"entity_type": "processing_activity", "entity_id": env["ids"]["proc_a"]}]
    r = await env["client"].post(DOCS, json=body, headers=who)
    assert r.status_code == 201, r.text
    return r.json()


async def _upload(env, who, doc_id, name, data, mime):
    r = await env["client"].post(f"{DOCS}/{doc_id}/versions", files={"file": (name, data, mime)}, headers=who)
    assert r.status_code == 200, r.text
    return r.json()


async def _events(env, event_type):
    async with env["maker"]() as db:
        return list((await db.execute(select(AuditLog).where(AuditLog.event_type == event_type))).scalars().all())


def _breach(unit="unit-a", **override):
    aware = (datetime.utcnow() - timedelta(hours=3)).replace(microsecond=0).isoformat()
    return {"title": "Invio errato", "organizational_unit_id": unit, "aware_at": aware, "description": "Testo riservato della violazione", **override}


# ================================================================ configurazione

def test_config_validation_keeps_admins_and_rejects_nonsense():
    ok = validate_config({"levels": {HR: {"roles": ["DPO"], "watermark": "required", "max_downloads_per_hour": 5}}})
    assert ok["levels"][HR]["roles"] == ["SUPER_ADMIN", "TENANT_ADMIN", "DPO"] and ok["levels"][HR]["watermark"] == "required"
    assert ok["step_up"] == {"max_age_seconds": 900, "admin_actions": False}
    for bad in ({"levels": {"INTERNAL": {}}}, {"levels": {HR: {"roles": ["CAPO"]}}}, {"levels": {HR: {"watermark": "sempre"}}},
                {"levels": {HR: {"max_downloads_per_hour": 0}}}, {"step_up": {"max_age_seconds": 10}}):
        with pytest.raises(ValueError):
            validate_config(bad)


def test_step_up_needs_strong_factors_and_a_recent_authentication():
    now = 1_000_000.0

    def user(**claims):
        return TokenPayload({"sub": "u", **claims})

    assert step_up_satisfied(user(amr=["pwd", "otp"], auth_time=now - 60), 900, now)
    assert step_up_satisfied(user(amr=["mfa"], iat=now - 60), 900, now)                       # senza auth_time si usa iat
    assert step_up_satisfied(user(acr="urn:example:mfa", amr=[], auth_time=now - 5), 900, now)
    assert not step_up_satisfied(user(amr=["pwd"], auth_time=now - 5), 900, now)             # un solo fattore debole
    assert not step_up_satisfied(user(amr=["pwd", "otp"], auth_time=now - 901), 900, now)    # troppo vecchia
    assert not step_up_satisfied(user(amr=["pwd", "otp"], auth_time=now + 500), 900, now)    # nel futuro
    assert not step_up_satisfied(user(amr=["otp"]), 900, now) and not step_up_satisfied(user(), 900, now)


@pytest.mark.asyncio
async def test_policy_endpoint_defaults_versioning_and_effect(env, w):
    c = env["client"]
    for who in ("mgr", "dpo", "op"):
        assert (await c.get(POLICY, headers=w[who])).status_code == 403                       # solo gli amministratori
    got = (await c.get(POLICY, headers=ADMIN)).json()
    assert got["row_version"] == 0 and got["levels"] == [RE, HR]
    assert got["config"]["levels"][HR]["block_bulk_export"] is True and got["config"]["levels"][HR]["watermark"] == "if_possible"
    assert got["config"]["levels"][RE]["roles"].count("OPERATOR") == 0 and got["config"]["step_up"]["admin_actions"] is False

    body = {"levels": {RE: {"roles": ["OPERATOR"], "watermark": "off"}}}
    assert (await c.put(POLICY, json=body, headers=ADMIN)).status_code == 428
    assert (await c.put(POLICY, json=body, headers=ifm(3, ADMIN))).status_code == 409
    assert (await c.put(POLICY, json=body, headers=ifm(0, w["mgr"]))).status_code == 403
    assert (await c.put(POLICY, json={"levels": {"INTERNAL": {}}}, headers=ifm(0, ADMIN))).status_code == 400
    assert (await c.put(POLICY, json={"levels": {HR: {"roles": ["CAPO"]}}}, headers=ifm(0, ADMIN))).status_code == 400
    assert (await c.put(POLICY, json={"levels": {HR: {"watermark": "sempre"}}}, headers=ifm(0, ADMIN))).status_code == 400
    saved = await _set_policy(env, body)
    assert saved["row_version"] == 1 and saved["config"]["levels"][RE]["roles"][-1] == "OPERATOR"
    assert "TENANT_ADMIN" in saved["config"]["levels"][RE]["roles"]                            # gli amministratori restano sempre
    assert saved["config"]["levels"][HR]["block_bulk_export"] is True                           # ciò che non si indica torna al default
    assert len(await _events(env, "policy.confidentiality.updated")) == 1


# ================================================================ livelli sugli elementi

@pytest.mark.asyncio
async def test_highly_restricted_applies_only_where_the_spec_says(env, w):
    c, ids = env["client"], env["ids"]
    proc = {"title": "Trattamento", "organizational_unit_id": "unit-a", "confidentiality_level": HR}
    assert (await c.post(f"{API}/processing-activities", json=proc, headers=w["mgr"])).status_code == 422
    assert (await c.post(f"{API}/assets", json={"name": "Sistema", "confidentiality_level": HR}, headers=w["mgr"])).status_code == 422
    assert (await c.post(f"{API}/assets", json={"name": "Sistema", "confidentiality_level": RE}, headers=w["mgr"])).status_code == 201
    doc = await _document(env, w["mgr"], HR)
    assert doc["confidentiality_level"] == HR
    assert (await c.post(BR, json=_breach(confidentiality_level=HR), headers=w["mgr"])).json()["confidentiality_level"] == HR
    assert (await c.post(DPIAS, json={"title": "DPIA", "processing_activity_id": ids["proc_a"], "confidentiality_level": HR}, headers=w["mgr"])).json()["confidentiality_level"] == HR
    project = {"title": "P", "proposing_unit_id": "unit-a", "project_owner_id": "ref-a", "confidentiality_level": HR}
    assert (await c.post(PR, json=project, headers=w["mgr"])).json()["confidentiality_level"] == HR


@pytest.mark.asyncio
async def test_documents_follow_the_level_in_lists_details_search_and_changes(env, w):
    c = env["client"]
    internal = await _document(env, w["mgr"], title="Circolare interna")
    restricted = await _document(env, w["mgr"], RE, title="Verbale riservato")
    secret = await _document(env, w["mgr"], HR, title="Parere legale segreto")
    assert internal["confidentiality_level"] == "INTERNAL"

    def titles(response):
        return sorted(d["title"] for d in response.json())

    everyone = ["Circolare interna", "Parere legale segreto", "Verbale riservato"]
    assert titles(await c.get(DOCS, headers=w["mgr"])) == everyone
    assert titles(await c.get(DOCS, headers=w["dpo"])) == everyone
    assert titles(await c.get(DOCS, headers=ADMIN)) == everyone
    assert titles(await c.get(DOCS, headers=w["ref-a"])) == ["Circolare interna", "Verbale riservato"]     # referente: riservato sì, altamente riservato no
    assert titles(await c.get(DOCS, headers=w["aud"])) == ["Circolare interna", "Verbale riservato"]
    assert titles(await c.get(DOCS, headers=w["reader"])) == ["Circolare interna"]                  # sola lettura: solo ciò che non è riservato
    for who in ("op", "reader", "ict"):
        assert "Parere legale segreto" not in titles(await c.get(DOCS, headers=w[who]))
        assert (await c.get(f"{DOCS}/{secret['id']}", headers=w[who])).status_code == 403
    assert (await c.get(f"{DOCS}/{restricted['id']}", headers=w["op"])).status_code == 403
    assert (await c.get(f"{DOCS}/{restricted['id']}", headers=w["ref-a"])).status_code == 200
    assert (await c.get(f"{DOCS}/{secret['id']}", headers=w["ref-a"])).status_code == 403
    assert (await c.get(f"{DOCS}/{secret['id']}", headers=w["b"])).status_code == 404

    # la ricerca globale non rivela titoli di ciò che non si può leggere
    found = (await c.get(f"{API}/search", params={"q": "legale"}, headers=w["ref-a"])).json()
    assert found["documents"] == [] and len((await c.get(f"{API}/search", params={"q": "legale"}, headers=w["dpo"])).json()["documents"]) == 1

    # chi non legge un livello non lo assegna né lo toglie
    assert (await c.post(DOCS, json={"title": "X", "confidentiality_level": HR}, headers=w["ref-a"])).status_code == 403
    assert (await c.post(DOCS, json={"title": "X", "confidentiality_level": RE}, headers=w["op"])).status_code == 403
    assert (await c.post(DOCS, json={"title": "Y", "confidentiality_level": RE}, headers=w["ref-a"])).status_code == 201
    url = f"{DOCS}/{restricted['id']}"
    body = {"title": "Verbale riservato", "confidentiality_level": HR}
    assert (await c.patch(url, json=body, headers=ifm(restricted["row_version"], w["ref-a"]))).status_code == 403      # non alza a un livello che non legge
    raised = await c.patch(url, json=body, headers=ifm(restricted["row_version"], w["dpo"]))
    assert raised.status_code == 200 and raised.json()["confidentiality_level"] == HR
    # una modifica che omette il livello non lo abbassa
    kept = await c.patch(url, json={"title": "Verbale riservato bis"}, headers=ifm(raised.json()["row_version"], w["dpo"]))
    assert kept.status_code == 200 and kept.json()["confidentiality_level"] == HR
    lowered = await c.patch(url, json={"title": "Verbale", "confidentiality_level": "INTERNAL"}, headers=ifm(kept.json()["row_version"], w["dpo"]))
    assert lowered.status_code == 200 and lowered.json()["confidentiality_level"] == "INTERNAL"


@pytest.mark.asyncio
async def test_author_and_configured_roles_keep_their_access(env, w):
    c = env["client"]
    mine = await _document(env, w["ref-a"], RE, title="Mio riservato")            # il referente legge RESTRICTED: può crearlo
    # l'ente restringe RESTRICTED al solo governo: l'autore continua a leggere ciò che ha creato
    await _set_policy(env, {"levels": {RE: {"roles": ["PRIVACY_MANAGER"]}}})
    assert (await c.get(f"{DOCS}/{mine['id']}", headers=w["ref-a"])).status_code == 200
    other = await _document(env, w["mgr"], RE, title="Del governo")
    assert (await c.get(f"{DOCS}/{other['id']}", headers=w["ref-a"])).status_code == 403
    assert (await c.get(f"{DOCS}/{other['id']}", headers=w["dpo"])).status_code == 403          # il DPO non è più tra i ruoli
    assert (await c.get(f"{DOCS}/{other['id']}", headers=w["mgr"])).status_code == 200
    assert (await c.get(f"{DOCS}/{other['id']}", headers=ADMIN)).status_code == 200            # gli amministratori sì, sempre
    await _set_policy(env, {"levels": {RE: {"roles": ["OPERATOR"]}}})
    assert (await c.get(f"{DOCS}/{other['id']}", headers=w["op"])).status_code == 200


@pytest.mark.asyncio
async def test_breaches_dpias_and_projects_follow_the_level_and_their_annotations_too(env, w):
    c, ids = env["client"], env["ids"]
    open_breach = (await c.post(BR, json=_breach(title="Aperta"), headers=w["mgr"])).json()
    secret = (await c.post(BR, json=_breach(title="Segreta", confidentiality_level=HR), headers=w["mgr"])).json()
    assert [b["title"] for b in (await c.get(BR, headers=w["ref-a"])).json()] == ["Aperta"]
    assert sorted(b["title"] for b in (await c.get(BR, headers=w["dpo"])).json()) == ["Aperta", "Segreta"]
    assert (await c.get(f"{BR}/{secret['id']}", headers=w["ref-a"])).status_code == 403
    assert (await c.get(f"{BR}/{secret['id']}", headers=w["ict"])).status_code == 403           # ICT gestisce le violazioni, ma non le altamente riservate
    assert (await c.get(f"{BR}/{secret['id']}", headers=w["dpo"])).status_code == 200
    # modificare senza indicare il livello non lo cambia; chi non lo legge non modifica
    upd = _breach(title="Segreta 2")
    assert (await c.patch(f"{BR}/{secret['id']}", json=upd, headers=ifm(secret["row_version"], w["ref-a"]))).status_code == 403
    done = await c.patch(f"{BR}/{secret['id']}", json=upd, headers=ifm(secret["row_version"], w["mgr"]))
    assert done.status_code == 200 and done.json()["confidentiality_level"] == HR
    assert (await c.post(BR, json=_breach(confidentiality_level=HR), headers=w["ref-a"])).status_code == 403

    # commenti, tag e campi: l'elemento non si vede, quindi nemmeno le sue annotazioni
    assert (await c.get(f"{API}/entities/data_breach/{secret['id']}/annotations", headers=w["ref-a"])).status_code == 403
    assert (await c.post(f"{API}/entities/data_breach/{secret['id']}/comments", json={"body": "x"}, headers=w["ref-a"])).status_code == 403
    assert (await c.get(f"{API}/entities/data_breach/{secret['id']}/annotations", headers=w["dpo"])).status_code == 200
    tag = (await c.post(f"{API}/tags", json={"name": "Sensibile"}, headers=w["mgr"])).json()
    for entity in (f"data_breach/{secret['id']}", f"data_breach/{open_breach['id']}"):
        assert (await c.put(f"{API}/entities/{entity}/tags", json={"tag_ids": [tag["id"]]}, headers=w["mgr"])).status_code == 200
    seen = (await c.get(f"{API}/tags/{tag['id']}/entities", headers=w["ref-a"])).json()
    assert [e["entity_id"] for e in seen] == [open_breach["id"]]

    # DPIA
    dpia = (await c.post(DPIAS, json={"title": "DPIA riservata", "processing_activity_id": ids["proc_a"], "confidentiality_level": HR}, headers=w["mgr"])).json()
    assert [d["id"] for d in (await c.get(DPIAS, headers=w["ref-a"])).json()] == []
    assert (await c.get(f"{DPIAS}/{dpia['id']}", headers=w["ref-a"])).status_code == 403
    assert (await c.get(f"{DPIAS}/{dpia['id']}", headers=w["dpo"])).status_code == 200
    assert (await c.get(f"{API}/search", params={"q": "riservata"}, headers=w["ref-a"])).json()["dpias"] == []
    plain = {"title": "DPIA riservata", "processing_activity_id": ids["proc_a"]}
    keep = await c.patch(f"{DPIAS}/{dpia['id']}", json=plain, headers=ifm(dpia["row_version"], w["mgr"]))
    assert keep.status_code == 200 and keep.json()["confidentiality_level"] == HR                # omesso = invariato

    # progetti: il responsabile legge il proprio, gli altri secondo il ruolo
    project = (await c.post(PR, json={"title": "Progetto sensibile", "description": "d", "proposing_unit_id": "unit-a", "project_owner_id": "ref-a",
                                      "confidentiality_level": HR}, headers=w["mgr"])).json()
    assert (await c.get(f"{PR}/{project['id']}", headers=w["ref-a"])).status_code == 200        # è il responsabile
    assert (await c.get(f"{PR}/{project['id']}", headers=w["op"])).status_code == 403
    assert [p["id"] for p in (await c.get(PR, headers=w["op"])).json()] == []
    assert [p["id"] for p in (await c.get(PR, headers=w["dpo"])).json()] == [project["id"]]


# ================================================================ download

@pytest.mark.asyncio
async def test_highly_restricted_downloads_are_watermarked_audited_and_leave_the_original_intact(env, w):
    c = env["client"]
    pdf, word = await _document(env, w["mgr"], HR, title="PDF"), await _document(env, w["mgr"], HR, title="Word")
    original = _pdf()
    await _upload(env, w["mgr"], pdf["id"], "parere.pdf", original, "application/pdf")
    await _upload(env, w["mgr"], word["id"], "parere.docx", _docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    r = await c.get(f"{DOCS}/{pdf['id']}/file", headers=w["dpo"])
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    text = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
    assert HR in text and pdf["code"] in text and "dpo" in text and "UTC" in text
    d = await c.get(f"{DOCS}/{word['id']}/file", headers=w["dpo"])
    assert d.status_code == 200
    header = " ".join(p.text for p in WordDocument(io.BytesIO(d.content)).sections[0].header.paragraphs)
    assert HR in header and word["code"] in header

    events = await _events(env, "document.download.highly_restricted")
    assert len(events) == 2 and {e.actor_user_id for e in events} == {"dpo"}
    meta = events[0].event_metadata
    assert meta["level"] == HR and meta["watermarked"] is True and meta["sha256"] and meta["watermark_mode"] == "if_possible"
    # il file archiviato non cambia: un secondo download parte dallo stesso originale
    again = await c.get(f"{DOCS}/{pdf['id']}/file", headers=w["mgr"])
    assert "mgr" in PdfReader(io.BytesIO(again.content)).pages[0].extract_text()

    # formato non marcabile: con «se possibile» si scarica intatto e l'audit lo dichiara; con «required» no
    note = await _document(env, w["mgr"], HR, title="Nota")
    await _upload(env, w["mgr"], note["id"], "nota.txt", b"ciao", "text/plain")
    plain = await c.get(f"{DOCS}/{note['id']}/file", headers=w["dpo"])
    assert plain.status_code == 200 and plain.content == b"ciao"
    assert (await _events(env, "document.download.highly_restricted"))[-1].event_metadata["watermarked"] is False
    await _set_policy(env, {"levels": {HR: {"watermark": "required"}}})
    blocked = await c.get(f"{DOCS}/{note['id']}/file", headers=w["dpo"])
    assert blocked.status_code == 409 and "filigrana" in blocked.json()["detail"]
    assert (await c.get(f"{DOCS}/{pdf['id']}/file", headers=w["dpo"])).status_code == 200
    await _set_policy(env, {"levels": {HR: {"watermark": "off"}}})
    assert (await c.get(f"{DOCS}/{pdf['id']}/file", headers=w["dpo"])).content == original          # senza filigrana: il file archiviato, intatto

    # chi non legge il livello non scarica; i documenti aperti si scaricano intatti e con l'evento di sempre
    assert (await c.get(f"{DOCS}/{pdf['id']}/file", headers=w["ref-a"])).status_code == 403
    open_doc = await _document(env, w["mgr"], title="Aperto")
    await _upload(env, w["mgr"], open_doc["id"], "a.txt", b"aperto", "text/plain")
    assert (await c.get(f"{DOCS}/{open_doc['id']}/file", headers=w["ref-a"])).content == b"aperto"
    assert len(await _events(env, "document.download")) == 1


@pytest.mark.asyncio
async def test_step_up_is_required_for_downloads_when_the_tenant_asks_for_it(env, w):
    c = env["client"]
    doc = await _document(env, w["mgr"], HR)
    await _upload(env, w["mgr"], doc["id"], "a.pdf", _pdf(), "application/pdf")
    url = f"{DOCS}/{doc['id']}/file"
    assert (await c.get(url, headers=w["dpo"])).status_code == 200                              # di default non è richiesto
    await _set_policy(env, {"levels": {HR: {"step_up": True, "watermark": "off"}}, "step_up": {"max_age_seconds": 600}})

    denied = await c.get(url, headers=w["dpo"])
    assert denied.status_code == 403 and denied.headers["x-step-up-required"] == "true" and denied.headers["x-step-up-max-age"] == "600"
    assert "MFA" in denied.json()["detail"]
    assert (await c.get(url, headers=strong("dpo", amr=("pwd",)))).status_code == 403          # un solo fattore
    assert (await c.get(url, headers=strong("dpo", age=3600))).status_code == 403              # autenticazione vecchia
    assert (await c.get(url, headers=strong("dpo", age=30))).status_code == 200
    assert (await c.get(url, headers=strong("dpo", age=30, key="iat"))).status_code == 200
    assert (await c.get(url, headers=strong("ref-a", age=1))).status_code == 403                # l'MFA non sostituisce il livello
    assert len(await _events(env, "document.download.highly_restricted")) == 3                  # i tentativi respinti non sono download


@pytest.mark.asyncio
async def test_hourly_download_cap_limits_mass_exfiltration(env, w):
    c = env["client"]
    await _set_policy(env, {"levels": {HR: {"max_downloads_per_hour": 2, "watermark": "off"}}})
    doc = await _document(env, w["mgr"], HR)
    await _upload(env, w["mgr"], doc["id"], "a.txt", b"x", "text/plain")
    url = f"{DOCS}/{doc['id']}/file"
    assert (await c.get(url, headers=w["dpo"])).status_code == 200 and (await c.get(url, headers=w["dpo"])).status_code == 200
    capped = await c.get(url, headers=w["dpo"])
    assert capped.status_code == 429 and "2 download all'ora" in capped.json()["detail"]
    assert (await c.get(url, headers=w["mgr"])).status_code == 200                              # il tetto è per utente
    async with env["maker"]() as db:                                                             # scaduta l'ora, si riparte
        for log in (await db.execute(select(AuditLog).where(AuditLog.event_type == "document.download.highly_restricted", AuditLog.actor_user_id == "dpo"))).scalars():
            log.event_time = datetime.utcnow() - timedelta(hours=2)
        await db.commit()
    assert (await c.get(url, headers=w["dpo"])).status_code == 200


# ================================================================ azioni amministrative sensibili e export

@pytest.mark.asyncio
async def test_sensitive_admin_actions_need_step_up_only_when_enabled(env, w):
    c = env["client"]
    async with env["maker"]() as db:                                                               # il catalogo dei ruoli non è precaricato nei test
        for role in (R.SUPER_ADMIN, R.OPERATOR):
            db.add(PrivacyRoleCatalog(code=role, name=role.value))
        await db.commit()
    assign = {"user_id": "nuovo", "role_code": "SUPER_ADMIN"}
    export = f"{API}/export/tenant"
    assert (await c.post(f"{API}/roles/assign", json=assign, headers=ADMIN)).status_code == 201        # di default: nessun controllo aggiuntivo
    await _set_policy(env, {"step_up": {"admin_actions": True, "max_age_seconds": 600}})

    for call in (lambda h: c.post(f"{API}/roles/assign", json=assign, headers=h), lambda h: c.post(export, headers=h)):
        denied = await call(ADMIN)
        assert denied.status_code == 403 and denied.headers["x-step-up-required"] == "true"
    admin_mfa = strong("root", tenant_role="admin", age=20)
    assert (await c.post(f"{API}/roles/assign", json={**assign, "user_id": "altro"}, headers=admin_mfa)).status_code == 201
    assert (await c.post(export, headers=admin_mfa)).status_code == 200
    # gli altri ruoli non richiedono nulla di più
    assert (await c.post(f"{API}/roles/assign", json={"user_id": "x1", "role_code": "OPERATOR"}, headers=ADMIN)).status_code == 201
    assert (await c.get(POLICY, headers=ADMIN)).status_code == 200


@pytest.mark.asyncio
async def test_retention_changes_need_step_up_when_enabled(env, w):
    c, pid = env["client"], env["ids"]["proc_a"]
    url = f"{API}/processing-activities/{pid}/retention-rules"
    rule = {"retention_value": 10, "retention_unit": "YEARS", "retention_criterion": "Prescrizione ordinaria",
            "start_trigger": "cessazione del rapporto", "legal_reference": "art. 2946 c.c.", "disposal_method": "DELETION"}
    assert (await c.post(url, json=rule, headers=w["mgr"])).status_code == 201                  # di default nessun controllo aggiuntivo
    await _set_policy(env, {"step_up": {"admin_actions": True}})
    denied = await c.post(url, json=rule, headers=w["mgr"])
    assert denied.status_code == 403 and denied.headers["x-step-up-required"] == "true"
    mfa = strong("mgr", age=5)
    assert (await c.post(url, json=rule, headers=mfa)).status_code == 201
    created = (await c.get(url, headers=w["mgr"])).json()[0]
    assert (await c.patch(f"{API}/retention-rules/{created['id']}", json=rule, headers=ifm(created["row_version"], w["mgr"]))).status_code == 403
    assert (await c.patch(f"{API}/retention-rules/{created['id']}", json=rule, headers=ifm(created["row_version"], mfa))).status_code == 200
    assert (await c.delete(f"{API}/retention-rules/{created['id']}", headers=w["mgr"])).status_code == 403
    assert (await c.delete(f"{API}/retention-rules/{created['id']}", headers=mfa)).status_code == 204


def _zip_json(content: bytes, name: str):
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        return json.loads(z.read(name))


@pytest.mark.asyncio
async def test_bulk_export_omits_the_content_of_elements_whose_level_blocks_it(env, w):
    c, ids = env["client"], env["ids"]
    secret = await _document(env, w["mgr"], HR, title="Documento segreto")
    open_doc = await _document(env, w["mgr"], title="Documento aperto")
    await _upload(env, w["mgr"], secret["id"], "segreto.txt", b"contenuto segreto", "text/plain")
    await _upload(env, w["mgr"], open_doc["id"], "aperto.txt", b"contenuto aperto", "text/plain")
    hidden = (await c.post(BR, json=_breach(title="Violazione segreta", confidentiality_level=HR), headers=w["mgr"])).json()
    await c.post(f"{API}/entities/data_breach/{hidden['id']}/comments", json={"body": "commento segreto"}, headers=w["mgr"])
    await c.post(BR, json=_breach(title="Violazione aperta"), headers=w["mgr"])
    await c.post(DPIAS, json={"title": "DPIA segreta", "processing_activity_id": ids["proc_a"], "confidentiality_level": HR}, headers=w["mgr"])

    r = await c.post(f"{API}/export/tenant", headers=ADMIN)
    assert r.status_code == 200
    raw = r.content
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = z.namelist()
        manifest = json.loads(z.read("manifest.json"))
        blob = b"".join(z.read(n) for n in names)
    assert not any("segreto.txt" in n for n in names) and any("aperto.txt" in n for n in names)
    for hidden_text in (b"Documento segreto", b"Violazione segreta", b"DPIA segreta", b"commento segreto", b"contenuto segreto"):
        assert hidden_text not in blob, hidden_text
    assert b"Documento aperto" in blob and b"Violazione aperta" in blob
    assert manifest["redacted"]["levels"] == [HR] and manifest["redacted"]["rows"]["privacypa_documents"] == 1
    assert manifest["redacted"]["rows"]["privacypa_data_breaches"] == 1 and manifest["redacted"]["rows"]["privacypa_dpias"] == 1
    assert manifest["redacted"]["rows"]["privacypa_comments"] == 1
    assert any(s["reason"].startswith("livello di riservatezza") for s in manifest["documents"]["skipped"])
    docs = _zip_json(raw, "data/documents.json")
    by_title = {d["code"]: d for d in docs}
    assert by_title[secret["code"]]["title"] is None and by_title[secret["code"]]["confidentiality_level"] == HR
    assert by_title[open_doc["code"]]["title"] == "Documento aperto"
    audit = {a["entity_id"]: a for a in _zip_json(raw, "data/audit_log.json") if a["entity_type"] == "document"}
    assert audit[secret["id"]]["new_values"] is None and audit[open_doc["id"]]["new_values"] is not None

    # l'ente può consentire l'export: allora c'è tutto
    await _set_policy(env, {"levels": {HR: {"block_bulk_export": False}}})
    full = await c.post(f"{API}/export/tenant", headers=ADMIN)
    with zipfile.ZipFile(io.BytesIO(full.content)) as z:
        assert any("segreto.txt" in n for n in z.namelist())
        assert json.loads(z.read("manifest.json"))["redacted"] == {"levels": [], "rows": {}}
        assert b"Documento segreto" in b"".join(z.read(n) for n in z.namelist())
