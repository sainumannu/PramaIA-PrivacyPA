"""Ricerca globale (§39) e query builder (§40)."""
import io
from datetime import datetime, timedelta

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import select

from backend.models.privacypa_audit import AuditLog
from backend.models.privacypa_catalog import DataCategoryCatalog, GdprDataClass
from backend.models.privacypa_identity import PrivacyRoleCode as R, ScopeType
from backend.models.privacypa_processing import ProcessingDataCategory
from backend.services.privacypa_query_service import ENTITIES, TYPE_OPERATORS
from backend.tests.conftest import TA, TB, _h, put_links

API = "/api/privacypa"
Q = f"{API}/queries"
ADMIN = _h(TA, "root", "admin")
HR = "HIGHLY_RESTRICTED"


def cond(field, operator, value=None):
    return {"field": field, "operator": operator, "value": value}


@pytest_asyncio.fixture
async def w(env):
    grant = env["grant"]
    for user, role, scope, unit in (
        ("mgr", R.PRIVACY_MANAGER, ScopeType.TENANT, None), ("dpo", R.DPO, ScopeType.TENANT, None),
        ("op", R.OPERATOR, ScopeType.TENANT, None), ("reader", R.READ_ONLY, ScopeType.TENANT, None),
        ("aud", R.AUDITOR, ScopeType.TENANT, None),
        ("ref-a", R.PRIVACY_REFERENT, ScopeType.UNIT, "unit-a"), ("ref-b", R.PRIVACY_REFERENT, ScopeType.UNIT, "unit-b"),
    ):
        await grant(TA, user, role, scope, unit)
    await grant(TB, "b-mgr", R.PRIVACY_MANAGER)
    return {k: _h(TA, k) for k in ("mgr", "dpo", "op", "reader", "aud", "ref-a", "ref-b")} | {"b": _h(TB, "b-mgr")}


async def _processing(env, title, unit="unit-a", **extra):
    r = await env["client"].post(f"{API}/processing-activities", json={"title": title, "organizational_unit_id": unit, **extra}, headers=_h(TA, "mgr"))
    assert r.status_code == 201, r.text
    return r.json()


async def _link_category(env, processing_id, code, gdpr_class):
    async with env["maker"]() as db:
        cat = DataCategoryCatalog(tenant_id=None, code=code, name=code.title(), gdpr_class=gdpr_class)
        db.add(cat)
        await db.flush()
        db.add(ProcessingDataCategory(processing_activity_id=processing_id, data_category_id=cat.id))
        await db.commit()


@pytest_asyncio.fixture
async def data(env, w):
    """Cinque trattamenti nell'ente A (due già presenti + tre nuovi) e uno nell'ente B; due con dati sanitari, uno con DPIA."""
    c = env["client"]
    video = await _processing(env, "Videosorveglianza urbana", next_review_date="2026-03-01", description="Telecamere in centro")
    health = await _processing(env, "Sanità digitale", profiling=True, next_review_date="2027-01-01")
    proto = await _processing(env, "Protocollo", unit="unit-b")
    await _link_category(env, video["id"], "SALUTE", GdprDataClass.SPECIAL_ART9)
    await _link_category(env, health["id"], "SALUTE", GdprDataClass.SPECIAL_ART9)
    await _link_category(env, proto["id"], "ANAGRAFICI", GdprDataClass.ORDINARY)
    dpia = await c.post(f"{API}/dpias", json={"title": "DPIA videosorveglianza", "processing_activity_id": video["id"]}, headers=w["mgr"])
    assert dpia.status_code == 201, dpia.text
    return {"video": video, "health": health, "proto": proto, "dpia": dpia.json()}


async def _run(env, who, entity, conditions=(), *, match="all", **extra):
    return await env["client"].post(f"{Q}/run", json={"entity": entity, "match": match, "conditions": list(conditions), **extra}, headers=who)


async def _titles(env, who, entity, conditions=(), **kw):
    r = await _run(env, who, entity, conditions, **kw)
    assert r.status_code == 200, r.text
    titles = [row["title"] for row in r.json()["rows"]]
    return titles if "sort" in kw else sorted(titles)             # senza ordinamento richiesto conta l'insieme, non l'ordine


PROC = "processing_activity"


# ================================================================ catalogo

def test_every_field_has_operators_and_columns_match_the_catalog():
    for key, entity in ENTITIES.items():
        assert "code" in entity.fields and "tag" in entity.fields, key
        assert set(entity.columns) <= {n for n, f in entity.fields.items() if f.expr is not None and f.multi is None and f.show}, key
        for name, f in entity.fields.items():
            assert f.operators == TYPE_OPERATORS[f.type], (key, name)
            if f.type == "enum":
                assert f.values, (key, name)
    assert "data_subject_request" not in ENTITIES                       # dati del richiedente: fuori dal query builder


@pytest.mark.asyncio
async def test_schema_shows_only_readable_modules_and_export_right(env, w):
    c = env["client"]
    full = (await c.get(f"{Q}/schema", headers=w["mgr"])).json()
    assert {e["key"] for e in full["entities"]} == set(ENTITIES) and full["can_export"] is True
    assert full["operators"] == ["equals", "not_equals", "contains", "in", "not_in", "date_before", "date_after", "is_null", "is_not_null"]
    proc = next(e for e in full["entities"] if e["key"] == PROC)
    fields = {f["key"]: f for f in proc["fields"]}
    assert fields["has_dpia"]["type"] == "flag" and fields["has_dpia"]["selectable"] is False
    assert fields["tag"]["multi"] is True and fields["status"]["values"] and fields["description"]["selectable"] is False
    assert fields["title"]["sortable"] is True and fields["tag"]["sortable"] is False

    op = (await c.get(f"{Q}/schema", headers=w["op"])).json()
    assert "data_breach" not in {e["key"] for e in op["entities"]} and "processing_activity" in {e["key"] for e in op["entities"]}
    assert op["can_export"] is False
    assert (await c.get(f"{Q}/schema", headers=w["aud"])).json()["can_export"] is True


# ================================================================ esecuzione

@pytest.mark.asyncio
async def test_spec_example_health_data_without_dpia(env, w, data):
    """§40: categoria dati = SALUTE AND DPIA = assente AND stato = ... (qui in bozza)."""
    got = await _titles(env, w["mgr"], PROC, [cond("data_category", "equals", "salute"), cond("has_dpia", "equals", False), cond("status", "equals", "DRAFT")])
    assert got == ["Sanità digitale"]
    assert await _titles(env, w["mgr"], PROC, [cond("data_category", "equals", "SALUTE"), cond("has_dpia", "equals", True)]) == ["Videosorveglianza urbana"]
    assert await _titles(env, w["mgr"], PROC, [cond("data_category", "equals", "SALUTE"), cond("has_dpia", "not_equals", True)]) == ["Sanità digitale"]
    assert await _titles(env, w["mgr"], PROC, [cond("data_class", "equals", "SPECIAL_ART9")]) == ["Sanità digitale", "Videosorveglianza urbana"]
    assert await _titles(env, w["mgr"], PROC, [cond("data_category", "not_equals", "SALUTE")]) == ["Protocollo", "Trattamento A", "Trattamento B"]
    assert await _titles(env, w["mgr"], PROC, [cond("data_category", "is_null")]) == ["Trattamento A", "Trattamento B"]
    assert await _titles(env, w["mgr"], PROC, [cond("dpia_status", "is_not_null")]) == ["Videosorveglianza urbana"]
    status = data["dpia"]["status"]
    assert await _titles(env, w["mgr"], PROC, [cond("dpia_status", "in", [status])]) == ["Videosorveglianza urbana"]
    assert await _titles(env, w["mgr"], PROC, [cond("has_processor", "equals", True)]) == []
    assert len(await _titles(env, w["mgr"], PROC, [cond("has_notice", "equals", False), cond("has_system", "equals", False)])) == 5


@pytest.mark.asyncio
async def test_operators_on_text_dates_booleans_lists_and_any(env, w, data):
    t = lambda *conds, **kw: _titles(env, w["mgr"], PROC, list(conds), **kw)
    assert await t(cond("title", "contains", "SANIT")) == ["Sanità digitale"]                          # senza distinguere maiuscole
    assert await t(cond("title", "equals", "protocollo")) == ["Protocollo"]
    assert await t(cond("title", "in", ["sanità digitale", "PROTOCOLLO"])) == ["Protocollo", "Sanità digitale"]
    assert await t(cond("title", "not_in", ["Sanità digitale", "Protocollo", "Trattamento A"])) == ["Trattamento B", "Videosorveglianza urbana"]
    assert len(await t(cond("title", "not_equals", "Protocollo"))) == 4
    assert await t(cond("title", "contains", "%")) == [] and await t(cond("title", "contains", "_")) == []          # i jolly sono caratteri normali
    assert await t(cond("next_review_date", "date_before", "2026-06-01")) == ["Videosorveglianza urbana"]
    assert await t(cond("next_review_date", "date_after", "2026-06-01")) == ["Sanità digitale"]
    assert await t(cond("next_review_date", "equals", "2026-03-01")) == ["Videosorveglianza urbana"]
    assert await t(cond("next_review_date", "date_after", "2026-03-01")) == ["Sanità digitale"]                    # «dopo» è stretto
    assert sorted(await t(cond("next_review_date", "is_null"))) == ["Protocollo", "Trattamento A", "Trattamento B"]
    assert len(await t(cond("next_review_date", "is_not_null"))) == 2
    assert await t(cond("profiling", "equals", True)) == ["Sanità digitale"] and len(await t(cond("profiling", "not_equals", True))) == 4
    assert await t(cond("description", "contains", "telecamere")) == ["Videosorveglianza urbana"]
    assert await t(cond("title", "contains", "sanit"), cond("title", "contains", "proto"), match="any") == ["Protocollo", "Sanità digitale"]
    assert await t(cond("title", "contains", "sanit"), cond("title", "contains", "proto"), match="all") == []
    assert len(await t()) == 5                                                                            # nessuna condizione: tutto ciò che si vede
    assert await t(cond("status", "in", ["DRAFT", "active"]), cond("organizational_unit_id", "equals", "unit-b")) == ["Protocollo", "Trattamento B"]


@pytest.mark.asyncio
async def test_datetime_number_and_enum_operators_on_breaches(env, w):
    c = env["client"]
    noon = datetime.utcnow().replace(hour=12, minute=0, second=0, microsecond=0)          # a mezzogiorno: il giorno non dipende dall'ora del test
    for title, when, subjects in (("Invio errato", noon - timedelta(days=10), 5), ("Smarrimento", noon - timedelta(days=2), 120), ("Accesso abusivo", noon - timedelta(days=1), None)):
        r = await c.post(f"{API}/breaches", json={"title": title, "organizational_unit_id": "unit-a", "aware_at": when.isoformat(),
                                                  "description": "Testo", "approx_data_subjects": subjects}, headers=w["mgr"])
        assert r.status_code == 201, r.text
    t = lambda *conds, **kw: _titles(env, w["mgr"], "data_breach", list(conds), **kw)
    day = (noon - timedelta(days=2)).date().isoformat()
    assert await t(cond("aware_at", "equals", day)) == ["Smarrimento"]                                       # un giorno intero, non un istante
    assert await t(cond("aware_at", "date_before", day)) == ["Invio errato"]
    assert await t(cond("aware_at", "date_after", day), sort="title") == ["Accesso abusivo"]
    assert await t(cond("approx_data_subjects", "equals", 5)) == ["Invio errato"]
    assert await t(cond("approx_data_subjects", "in", [5, "120"])) == ["Invio errato", "Smarrimento"]
    assert await t(cond("approx_data_subjects", "is_null")) == ["Accesso abusivo"]
    assert await t(cond("approx_data_subjects", "not_equals", 5)) == ["Accesso abusivo", "Smarrimento"]        # i vuoti contano come «diversi»
    assert await t(cond("risk_level", "is_null")) == ["Accesso abusivo", "Invio errato", "Smarrimento"]
    assert len(await t(cond("status", "not_in", ["CLOSED"]))) == 3


@pytest.mark.asyncio
async def test_sorting_pagination_columns_and_row_shape(env, w, data):
    c = env["client"]
    page = (await _run(env, w["mgr"], PROC, [], sort="title", direction="desc", limit=2, offset=1)).json()
    assert page["total"] == 5 and page["offset"] == 1 and page["limit"] == 2
    assert [r["title"] for r in page["rows"]] == ["Trattamento B", "Trattamento A"]
    assert [col["key"] for col in page["columns"]] == ["code", "title", "status", "organizational_unit_id", "next_review_date"]
    row = (await _run(env, w["mgr"], PROC, [cond("title", "equals", "Sanità digitale")], columns=["title", "profiling", "next_review_date"])).json()["rows"][0]
    assert row["values"] == {"title": "Sanità digitale", "profiling": True, "next_review_date": "2027-01-01"} and row["code"]
    assert (await _run(env, w["mgr"], PROC, [], limit=0)).status_code == 422 and (await _run(env, w["mgr"], PROC, [], limit=500)).status_code == 422
    assert (await _run(env, w["mgr"], PROC, [], offset=-1)).status_code == 422
    assert (await _run(env, w["mgr"], PROC, [], offset=99)).json()["rows"] == []
    assert (await c.get(f"{Q}/run", headers=w["mgr"])).status_code == 405


@pytest.mark.asyncio
async def test_invalid_queries_are_rejected_with_a_reason(env, w, data):
    bad = [
        ({"entity": "widget"}, "non interrogabile"),
        ({"entity": "data_subject_request"}, "non interrogabile"),
        ({"entity": PROC, "conditions": [cond("password", "equals", "x")]}, "Campo sconosciuto"),
        ({"entity": PROC, "conditions": [cond("title", "greater", "x")]}, "Operatore"),
        ({"entity": PROC, "conditions": [cond("status", "contains", "DR")]}, "Operatore"),
        ({"entity": PROC, "conditions": [cond("status", "equals", "BOZZA")]}, "Valore non ammesso"),
        ({"entity": PROC, "conditions": [cond("next_review_date", "date_before", "ieri")]}, "data"),
        ({"entity": PROC, "conditions": [cond("title", "equals", "")]}, "testo non vuoto"),
        ({"entity": PROC, "conditions": [cond("title", "equals")]}, "manca il valore"),
        ({"entity": PROC, "conditions": [cond("title", "equals", "x" * 300)]}, "255"),
        ({"entity": PROC, "conditions": [cond("title", "in", "Protocollo")]}, "elenco"),
        ({"entity": PROC, "conditions": [cond("title", "in", [])]}, "elenco"),
        ({"entity": PROC, "conditions": [cond("title", "in", [str(i) for i in range(51)])]}, "elenco"),
        ({"entity": PROC, "conditions": [cond("profiling", "equals", "forse")]}, "sì/no"),
        ({"entity": "data_breach", "conditions": [cond("approx_data_subjects", "equals", True)]}, "intero"),
        ({"entity": "data_breach", "conditions": [cond("approx_data_subjects", "equals", "molti")]}, "intero"),
        ({"entity": PROC, "conditions": [cond("has_dpia", "is_null")]}, "Operatore"),
        ({"entity": PROC, "conditions": [cond("title", "equals", "x")] * 21}, None),
        ({"entity": PROC, "match": "xor"}, None),
        ({"entity": PROC, "sort": "tag"}, "Ordinamento"),
        ({"entity": PROC, "sort": "nonesiste"}, "Ordinamento"),
        ({"entity": PROC, "direction": "up"}, None),
        ({"entity": PROC, "columns": ["description"]}, "Colonna"),
        ({"entity": PROC, "columns": ["has_dpia"]}, "Colonna"),
        ({"entity": PROC, "columns": ["tag"]}, "Colonna"),
    ]
    for body, reason in bad:
        r = await env["client"].post(f"{Q}/run", json=body, headers=w["mgr"])
        assert r.status_code == 422, (body, r.text)
        if reason:
            assert reason in r.text, (body, r.text)
    # un attacco a colpi di SQL è solo un testo da cercare
    r = await _run(env, w["mgr"], PROC, [cond("title", "equals", "x'; DROP TABLE privacypa_processing_activities; --")])
    assert r.status_code == 200 and r.json()["total"] == 0
    assert (await _run(env, w["mgr"], PROC)).json()["total"] == 5


# ================================================================ visibilità

@pytest.mark.asyncio
async def test_results_and_counts_only_include_what_the_user_may_see(env, w, data):
    ref_a, ref_b = w["ref-a"], w["ref-b"]
    for who, expected in ((ref_a, ["Sanità digitale", "Trattamento A", "Videosorveglianza urbana"]), (ref_b, ["Protocollo", "Trattamento B"])):
        r = (await _run(env, who, PROC)).json()
        assert sorted(x["title"] for x in r["rows"]) == expected and r["total"] == len(expected)          # anche il conteggio
    # una condizione non permette di dedurre l'esistenza di ciò che non si vede
    assert (await _run(env, ref_b, PROC, [cond("data_category", "equals", "SALUTE")])).json()["total"] == 0
    assert (await _run(env, ref_b, PROC, [cond("title", "contains", "Sanit")])).json()["rows"] == []
    assert await _titles(env, w["b"], PROC) == ["Trattamento X"]                                          # altro ente
    assert (await _run(env, w["b"], PROC, [cond("title", "contains", "Sanit")])).json()["total"] == 0
    # DPIA: legata al trattamento, quindi alla struttura
    assert (await _run(env, ref_a, "dpia")).json()["total"] == 1 and (await _run(env, ref_b, "dpia")).json()["total"] == 0


@pytest.mark.asyncio
async def test_module_permission_and_confidentiality_apply_to_queries(env, w):
    c = env["client"]
    aware = (datetime.utcnow() - timedelta(hours=3)).replace(microsecond=0).isoformat()
    for title, unit, level in (("Aperta", "unit-a", None), ("Segreta", "unit-a", HR), ("Altra struttura", "unit-b", None)):
        body = {"title": title, "organizational_unit_id": unit, "aware_at": aware, "description": "Testo"} | ({"confidentiality_level": level} if level else {})
        assert (await c.post(f"{API}/breaches", json=body, headers=w["mgr"])).status_code == 201
    for who in ("op", "reader"):
        assert (await _run(env, w[who], "data_breach")).status_code == 403                                   # modulo non consultabile
    assert await _titles(env, w["mgr"], "data_breach", sort="title") == ["Altra struttura", "Aperta", "Segreta"]
    assert await _titles(env, w["ref-a"], "data_breach") == ["Aperta"]                                        # struttura e livello
    assert (await _run(env, w["ref-a"], "data_breach")).json()["total"] == 1
    assert (await _run(env, w["ref-a"], "data_breach", [cond("confidentiality_level", "equals", HR)])).json()["total"] == 0
    assert (await _run(env, w["dpo"], "data_breach", [cond("confidentiality_level", "equals", HR)])).json()["total"] == 1


@pytest.mark.asyncio
async def test_queries_by_tag_and_tags_in_the_results(env, w, data):
    c = env["client"]
    tags = {}
    for name in ("PNRR", "Urgente"):
        tags[name] = (await c.post(f"{API}/tags", json={"name": name}, headers=w["mgr"])).json()["id"]
    put = lambda pid, *names: c.put(f"{API}/entities/{PROC}/{pid}/tags", json={"tag_ids": [tags[n] for n in names]}, headers=w["mgr"])
    assert (await put(data["video"]["id"], "PNRR", "Urgente")).status_code == 200
    assert (await put(data["health"]["id"], "PNRR")).status_code == 200
    t = lambda *conds: _titles(env, w["mgr"], PROC, list(conds))
    assert await t(cond("tag", "equals", "pnrr")) == ["Sanità digitale", "Videosorveglianza urbana"]
    assert await t(cond("tag", "equals", "PNRR"), cond("tag", "equals", "urgente")) == ["Videosorveglianza urbana"]      # due tag distinti sullo stesso elemento
    assert await t(cond("tag", "contains", "urg")) == ["Videosorveglianza urbana"]
    assert await t(cond("tag", "not_in", ["PNRR"])) == ["Protocollo", "Trattamento A", "Trattamento B"]
    assert await t(cond("tag", "is_null")) == ["Protocollo", "Trattamento A", "Trattamento B"]
    assert len(await t(cond("tag", "is_not_null"))) == 2
    rows = (await _run(env, w["mgr"], PROC, [cond("tag", "is_not_null")], sort="title")).json()["rows"]
    assert [r["tags"] for r in rows] == [["PNRR"], ["PNRR", "Urgente"]]
    assert (await _run(env, w["ref-b"], PROC, [cond("tag", "equals", "PNRR")])).json()["total"] == 0


@pytest.mark.asyncio
async def test_other_entities_have_queryable_fields(env, w):
    c, mgr = env["client"], w["mgr"]
    supplier = (await c.post(f"{API}/parties", json={"legal_name": "XYZ S.r.l.", "party_type": "PROCESSOR", "country_code": "IT", "tax_id": "01234567897", "processor_profile": {"appointment_end": "2099-12-31"}}, headers=mgr)).json()
    asset = (await c.post(f"{API}/assets", json={"name": "VMS", "asset_type": "APPLICATION", "cloud_service": True}, headers=mgr)).json()
    await c.post(f"{API}/assets", json={"name": "Portale", "asset_type": "APPLICATION"}, headers=mgr)
    project = {"title": "Nuovo servizio", "proposing_unit_id": "unit-a", "project_owner_id": "ref-a", "uses_cloud": True}
    assert (await c.post(f"{API}/projects", json=project, headers=mgr)).status_code == 201
    assert (await c.post(f"{API}/projects", json={**project, "title": "Altro", "uses_cloud": False}, headers=mgr)).status_code == 201
    assert await _titles(env, mgr, "party", [cond("country_code", "equals", "it")]) == ["XYZ S.r.l."]
    assert await _titles(env, mgr, "party", [cond("party_type", "equals", "PROCESSOR")]) == ["XYZ S.r.l."]
    assert await _titles(env, mgr, "asset", [cond("cloud_service", "equals", True)]) == ["VMS"]
    assert await _titles(env, mgr, "asset", [cond("has_processing", "equals", False)], sort="name") == ["Portale", "VMS"]
    assert asset["name"] == "VMS" and supplier["legal_name"] == "XYZ S.r.l."
    assert await _titles(env, mgr, "project", [cond("uses_cloud", "equals", True)]) == ["Nuovo servizio"]
    assert await _titles(env, mgr, "project", [cond("status", "equals", "PROPOSAL")], sort="title") == ["Altro", "Nuovo servizio"]
    for entity in ("dpia", "contract", "privacy_notice", "corrective_action", "audit"):
        assert (await _run(env, mgr, entity)).status_code == 200, entity


@pytest.mark.asyncio
async def test_link_flags_see_systems_processors_and_notices(env, w, data):
    c, mgr = env["client"], w["mgr"]
    asset = (await c.post(f"{API}/assets", json={"name": "VMS", "asset_type": "APPLICATION"}, headers=mgr)).json()
    party = (await c.post(f"{API}/parties", json={"legal_name": "XYZ S.r.l.", "party_type": "PROCESSOR", "tax_id": "01234567897", "processor_profile": {"appointment_end": "2099-12-31"}}, headers=mgr)).json()
    notice = await c.post(f"{API}/notices", json={"title": "Informativa videosorveglianza", "owner_unit_id": "unit-a"}, headers=mgr)
    assert notice.status_code == 201, notice.text
    proc_url = f"{API}/processing-activities/{data['video']['id']}"
    assert (await put_links(c, f"{proc_url}/assets", [{"asset_id": asset["id"], "role": "PROCESSING"}], mgr)).status_code == 200
    assert (await put_links(c, f"{proc_url}/processors", [{"party_id": party["id"]}], mgr)).status_code == 200
    assert (await put_links(c, f"{API}/notices/{notice.json()['id']}/processing-links", {"processing_activity_ids": [data["video"]["id"]]}, mgr)).status_code == 200
    only_video = ["Videosorveglianza urbana"]
    for flag in ("has_system", "has_processor", "has_notice"):
        assert await _titles(env, mgr, PROC, [cond(flag, "equals", True)]) == only_video, flag
        assert len(await _titles(env, mgr, PROC, [cond(flag, "equals", False)])) == 4, flag
    assert await _titles(env, mgr, "asset", [cond("has_processing", "equals", True)]) == ["VMS"]
    assert await _titles(env, mgr, "privacy_notice", [cond("has_processing", "equals", True)]) == ["Informativa videosorveglianza"]
    assert await _titles(env, mgr, "privacy_notice", [cond("has_processing", "equals", False)]) == []


# ================================================================ esportazione

async def _export(env, who, entity, conditions=(), fmt="csv", **extra):
    return await env["client"].post(f"{Q}/export", json={"entity": entity, "conditions": list(conditions), "format": fmt, **extra}, headers=who)


@pytest.mark.asyncio
async def test_export_needs_its_own_permission(env, w, data):
    for who in ("op", "reader", "ref-a", "dpo"):
        expected = 200 if who == "dpo" else 403
        assert (await _export(env, w[who], PROC)).status_code == expected, who
    assert (await _export(env, w["aud"], PROC)).status_code == 200
    assert (await _export(env, w["op"], PROC, [cond("nonesiste", "equals", "x")])).status_code == 403            # prima il permesso, poi il resto
    assert (await _export(env, w["mgr"], "data_subject_request")).status_code == 422
    assert (await _export(env, w["mgr"], PROC, fmt="pdf")).status_code == 422
    assert (await _export(env, w["mgr"], PROC, [cond("nonesiste", "equals", "x")])).status_code == 422
    # chi non può leggere il modulo non lo esporta nemmeno con il permesso di esportazione
    async with env["maker"]() as db:
        from backend.models.privacypa_approvals import TenantPolicy
        db.add(TenantPolicy(tenant_id=TA, permission_overrides={"breach.view": ["DPO"]}))
        await db.commit()
    assert (await _export(env, w["aud"], "data_breach")).status_code == 403


@pytest.mark.asyncio
async def test_csv_export_content_scope_and_formula_neutralization(env, w, data):
    await _processing(env, "=HYPERLINK(\"http://x\")", unit="unit-a")
    r = await _export(env, w["mgr"], PROC, [cond("organizational_unit_id", "equals", "unit-a")], columns=["title", "next_review_date"], sort="title")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"] == 'attachment; filename="interrogazione-processing_activity.csv"' and r.headers["x-excluded-rows"] == "0"
    lines = r.text.strip().splitlines()
    assert lines[0] == "title;next_review_date;tag"
    assert "'=HYPERLINK" in r.text and "\n=HYPERLINK" not in r.text
    assert "Sanità digitale;2027-01-01;" in lines and "Protocollo" not in r.text                      # solo unit-a


@pytest.mark.asyncio
async def test_xlsx_export_has_information_sheet_and_audit_without_values(env, w, data):
    r = await _export(env, w["mgr"], PROC, [cond("title", "contains", "sanit"), cond("has_dpia", "equals", False)], fmt="xlsx", columns=["code", "title"])
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Risultati", "Informazioni"]
    rows = list(book["Risultati"].iter_rows(values_only=True))
    assert rows[0] == ("code", "title", "tag") and rows[1][1] == "Sanità digitale" and len(rows) == 2
    info = {row[0]: row[1] for row in book["Informazioni"].iter_rows(values_only=True)}
    assert info["Ente"] == TA and info["Utente"] == "mgr@example.org" and info["Righe"] == "1" and "Estratto il" in info
    assert "title contains sanit" in info["  1"]
    async with env["maker"]() as db:
        events = list((await db.execute(select(AuditLog).where(AuditLog.event_type == "query.export"))).scalars().all())
    assert len(events) == 1
    meta = events[0].event_metadata
    assert meta["entity"] == PROC and meta["format"] == "xlsx" and meta["rows"] == 1 and meta["excluded_by_confidentiality"] == 0
    assert meta["conditions"] == [{"field": "title", "operator": "contains"}, {"field": "has_dpia", "operator": "equals"}]
    assert "sanit" not in str(events[0].event_metadata).lower()                                         # niente valori cercati


@pytest.mark.asyncio
async def test_bulk_export_skips_levels_that_block_it(env, w):
    c = env["client"]
    aware = (datetime.utcnow() - timedelta(hours=3)).replace(microsecond=0).isoformat()
    for title, level in (("Aperta", None), ("Segreta", HR)):
        body = {"title": title, "organizational_unit_id": "unit-a", "aware_at": aware, "description": "Testo"} | ({"confidentiality_level": level} if level else {})
        assert (await c.post(f"{API}/breaches", json=body, headers=w["mgr"])).status_code == 201
    assert (await _run(env, w["mgr"], "data_breach")).json()["total"] == 2                                # la si vede...
    r = await _export(env, w["mgr"], "data_breach", columns=["title"])
    assert r.text.strip().splitlines() == ["title;tag", "Aperta;"] and r.headers["x-excluded-rows"] == "1"      # ...ma non si esporta in massa
    async with env["maker"]() as db:
        event = (await db.execute(select(AuditLog).where(AuditLog.event_type == "query.export"))).scalars().one()
    assert event.event_metadata["excluded_by_confidentiality"] == 1 and event.event_metadata["rows"] == 1


# ================================================================ ricerca globale

@pytest.mark.asyncio
async def test_search_covers_description_tags_and_new_domains(env, w, data):
    c = env["client"]
    tag = (await c.post(f"{API}/tags", json={"name": "Videocamere"}, headers=w["mgr"])).json()["id"]
    assert (await c.put(f"{API}/entities/{PROC}/{data['health']['id']}/tags", json={"tag_ids": [tag]}, headers=w["mgr"])).status_code == 200
    aware = (datetime.utcnow() - timedelta(hours=3)).replace(microsecond=0).isoformat()
    assert (await c.post(f"{API}/breaches", json={"title": "Perdita chiavetta", "organizational_unit_id": "unit-a", "aware_at": aware,
                                                   "description": "Chiavetta con elenco videocamere"}, headers=w["mgr"])).status_code == 201
    assert (await c.post(f"{API}/projects", json={"title": "Videocamere nuove", "proposing_unit_id": "unit-a", "project_owner_id": "ref-a"}, headers=w["mgr"])).status_code == 201

    s = lambda who, q: c.get(f"{API}/search", params={"q": q}, headers=who)
    found = (await s(w["mgr"], "telecamere")).json()
    assert [p["title"] for p in found["processing_activities"]] == ["Videosorveglianza urbana"]            # descrizione
    found = (await s(w["mgr"], "videocamere")).json()
    assert [p["title"] for p in found["processing_activities"]] == ["Sanità digitale"]                    # tag
    assert [p["title"] for p in found["data_breaches"]] == ["Perdita chiavetta"]                          # descrizione della violazione
    assert [p["title"] for p in found["projects"]] == ["Videocamere nuove"]
    assert found["counts"] == {"processing_activities": 1, "data_breaches": 1, "projects": 1} and found["total"] == 3
    assert {"assets", "parties", "documents", "tasks", "dpias", "notices", "transfers", "risks", "contracts", "corrective_actions", "audits"} <= set(found)
    dpia = (await s(w["mgr"], "DPIA vide")).json()
    assert [d["title"] for d in dpia["dpias"]] == ["DPIA videosorveglianza"] and dpia["dpias"][0]["status"]
    # modulo non consultabile, struttura e altro ente: non compaiono né nei risultati né nei conteggi
    op = (await s(w["op"], "videocamere")).json()
    assert op["data_breaches"] == [] and "data_breaches" not in op["counts"] and op["total"] == 2
    assert (await s(w["ref-b"], "videocamere")).json()["total"] == 0
    assert (await s(w["b"], "videocamere")).json()["total"] == 0
    assert (await s(w["ref-a"], "videocamere")).json()["total"] == 3


@pytest.mark.asyncio
async def test_search_treats_wildcards_literally_and_respects_levels(env, w, data):
    c = env["client"]
    for q in ("%%", "__", "\\\\"):
        assert (await c.get(f"{API}/search", params={"q": q}, headers=w["mgr"])).json()["total"] == 0, q
    assert (await c.get(f"{API}/search", params={"q": "a"}, headers=w["mgr"])).status_code == 422             # minimo due caratteri
    aware = (datetime.utcnow() - timedelta(hours=3)).replace(microsecond=0).isoformat()
    body = {"title": "Violazione segreta", "organizational_unit_id": "unit-a", "aware_at": aware, "description": "Testo", "confidentiality_level": HR}
    assert (await c.post(f"{API}/breaches", json=body, headers=w["mgr"])).status_code == 201
    dpia_body = {"title": "DPIA segreta", "processing_activity_id": data["video"]["id"], "confidentiality_level": HR}
    assert (await c.post(f"{API}/dpias", json=dpia_body, headers=w["mgr"])).status_code == 201
    hidden = (await c.get(f"{API}/search", params={"q": "segreta"}, headers=w["ref-a"])).json()
    assert hidden["total"] == 0 and hidden["counts"] == {}
    seen = (await c.get(f"{API}/search", params={"q": "segreta"}, headers=w["dpo"])).json()
    assert seen["counts"] == {"data_breaches": 1, "dpias": 1}
