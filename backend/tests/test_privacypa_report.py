"""Report Registro art. 30 (§41, §153): nomi leggibili, titolare/DPO, scope, sicurezza CSV, audit export."""
import pytest
import pytest_asyncio

from backend.models.privacypa_assets import Asset
from backend.models.privacypa_catalog import (
    DataCategoryCatalog,
    DataSubjectCategoryCatalog,
    GdprDataClass,
    LegalBasisCatalog,
)
from backend.models.privacypa_identity import PrivacyRoleCode as R, ScopeType
from backend.models.privacypa_organization import Organization
from backend.models.privacypa_parties import Party, PartyType
from backend.services.privacypa_audit_service import AuditService
from backend.services.privacypa_processing_service import ProcessingService
from backend.tests.conftest import put_links, TA, _h

URL = "/api/privacypa/reports/art30"


@pytest_asyncio.fixture
async def data(env):
    """Trattamento completo nella struttura A, con anagrafiche referenziate per nome."""
    maker, ids = env["maker"], env["ids"]
    async with maker() as db:
        org = await db.get(Organization, f"org-{TA}")
        org.name, org.dpo_name, org.dpo_contact = "Comune di Prova", "Mario DPO", "dpo@comune.example"
        subj = DataSubjectCategoryCatalog(tenant_id=None, code="CIT", name="Cittadini")
        cat_ord = DataCategoryCatalog(tenant_id=None, code="ANAG", name="Dati anagrafici", gdpr_class=GdprDataClass.ORDINARY)
        cat_sp = DataCategoryCatalog(tenant_id=None, code="SAL", name="Dati sulla salute", gdpr_class=GdprDataClass.SPECIAL_ART9)
        basis = LegalBasisCatalog(code="GDPR_6_1_E", article="6", paragraph="1", letter="e", description="Compito di interesse pubblico")
        party = Party(tenant_id=TA, code="FO-000001", legal_name="Fornitore Beta S.r.l.", party_type=PartyType.PROCESSOR)
        asset = Asset(tenant_id=TA, code="AS-000001", name="Gestionale Servizi Sociali")
        db.add_all([subj, cat_ord, cat_sp, basis, party, asset])
        await db.flush()

        svc = ProcessingService(db)
        await svc.update_activity(
            TA, ids["proc_a"], "seed",
            title="=HYPERLINK(\"http://evil\",\"clic\")",
            organizational_unit_id="unit-a", description="Gestione servizi sociali",
            purposes=["Istruttoria prestazioni"], legal_basis_codes=["GDPR_6_1_E"],
            data_subject_ids=[subj.id], data_category_ids=[cat_ord.id, cat_sp.id],
        )
        await db.commit()
        return {"party": party.id, "asset": asset.id, "subject": subj.id, "cat": cat_sp.id}


async def _link(env, data):
    await env["grant"](TA, "mgr", R.PRIVACY_MANAGER)
    who = _h(TA, "mgr")
    base = f"/api/privacypa/processing-activities/{env['ids']['proc_a']}"
    assert (await put_links(env["client"], f"{base}/assets", json=[{"asset_id": data["asset"], "role": "STORAGE"}], headers=who)).status_code == 200
    assert (await put_links(env["client"], f"{base}/recipients", json=[{"party_id": data["party"], "communication_type": "COMMUNICATION"}], headers=who)).status_code == 200
    assert (await put_links(env["client"], f"{base}/processors", json=[{"party_id": data["party"]}], headers=who)).status_code == 200
    return who


@pytest.mark.asyncio
async def test_report_resolves_names_and_declares_controller(env, data):
    who = await _link(env, data)
    r = await env["client"].get(URL, headers=who)
    assert r.status_code == 200
    body = r.json()
    assert body["schema_version"] == "2.0"
    assert body["controller"]["name"] == "Comune di Prova"
    assert body["controller"]["dpo_name"] == "Mario DPO"
    assert body["unavailable_sections"] == []

    rec = next(x for x in body["processing_activities"] if x["id"] == env["ids"]["proc_a"])
    assert rec["unit_name"] == "unit-a"
    assert rec["data_subject_names"] == ["Cittadini"]
    assert {c["name"] for c in rec["data_categories_resolved"]} == {"Dati anagrafici", "Dati sulla salute"}
    assert rec["has_special_categories"] is True
    assert rec["legal_bases_resolved"][0]["description"] == "Compito di interesse pubblico"
    assert rec["recipients_resolved"] == [{"name": "Fornitore Beta S.r.l.", "communication_type": "COMMUNICATION"}]
    assert rec["processors_resolved"] == [{"name": "Fornitore Beta S.r.l."}]
    assert rec["systems"] == [{"name": "Gestionale Servizi Sociali", "role": "STORAGE"}]


@pytest.mark.asyncio
async def test_csv_uses_names_not_ids_and_neutralizes_formulas(env, data):
    who = await _link(env, data)
    r = await env["client"].get(URL, params={"format": "csv"}, headers=who)
    text = r.text
    for expected in ("Cittadini", "Dati sulla salute", "Fornitore Beta S.r.l.", "Gestionale Servizi Sociali",
                     "Comune di Prova", "Mario DPO", "conservazione"):
        assert expected in text
    assert data["subject"] not in text and data["cat"] not in text
    # il titolo che inizia con "=" non deve essere interpretabile come formula
    assert "'=HYPERLINK" in text and ";=HYPERLINK" not in text and "\n=HYPERLINK" not in text


@pytest.mark.asyncio
async def test_pdf_contains_readable_names_and_wraps_long_lines(env, data):
    who = await _link(env, data)
    r = await env["client"].get(URL, params={"format": "pdf"}, headers=who)
    assert r.content.startswith(b"%PDF-")
    for expected in (b"Cittadini", b"Fornitore Beta", b"Comune di Prova", b"BOZZA"):
        assert expected in r.content


@pytest.mark.asyncio
async def test_report_respects_scope_and_audits_export(env, data):
    await env["grant"](TA, "ref-a", R.PRIVACY_REFERENT, ScopeType.UNIT, "unit-a")
    r = await env["client"].get(URL, headers=_h(TA, "ref-a"))
    body = r.json()
    assert [x["id"] for x in body["processing_activities"]] == [env["ids"]["proc_a"]]
    assert body["filters"]["scope_limited"] is True

    async with env["maker"]() as db:
        events = await AuditService(db).list_recent(TA, entity_type="report")
    assert events and events[0].event_type == "report.art30.export"
    assert events[0].event_metadata["scope_limited"] is True


@pytest.mark.asyncio
async def test_xlsx_export_is_a_readable_workbook_without_formulas(env, data):
    import io

    from openpyxl import load_workbook

    who = await _link(env, data)
    r = await env["client"].get(URL, params={"format": "xlsx"}, headers=who)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Registro", "Conservazione", "Informazioni"]

    reg = wb["Registro"]
    header = [c.value for c in reg[1]]
    row = next(r_ for r_ in reg.iter_rows(min_row=2) if r_[0].value == "TR-000001")
    cells = dict(zip(header, row))
    assert cells["Struttura"].value == "unit-a" and "Fornitore Beta S.r.l." in cells["Responsabili"].value
    assert cells["Categorie particolari"].value == "sì" and cells["Conservazione"].value == "non definita"
    # il titolo che inizia con "=" resta testo, mai una formula
    assert cells["Titolo"].data_type == "s" and not str(cells["Titolo"].value).startswith("=")
    assert all(c.data_type != "f" for line in reg.iter_rows() for c in line)

    info = {r_[0].value: r_[1].value for r_ in wb["Informazioni"].iter_rows(min_row=2)}
    assert info["Titolare"] == "Comune di Prova" and info["DPO"] == "Mario DPO" and info["Trattamenti"] == 2
    audit = (await env["client"].get("/api/privacypa/audit", params={"entity_type": "report"}, headers=_h(TA, "root", "admin"))).json()
    assert audit[0]["metadata"]["format"] == "xlsx"


@pytest.mark.asyncio
async def test_docx_export_is_a_valid_document_with_readable_names(env, data):
    import io

    from docx import Document as WordDocument

    who = await _link(env, data)
    r = await env["client"].get(URL, params={"format": "docx"}, headers=who)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    doc = WordDocument(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Registro dei trattamenti (art. 30 GDPR)" in text and "BOZZA" in text
    for expected in ("Comune di Prova", "Mario DPO", "Cittadini", "Fornitore Beta S.r.l.", "Gestionale Servizi Sociali", "non definita"):
        assert expected in cells, expected
    assert data["subject"] not in " ".join(cells)
    audit = (await env["client"].get("/api/privacypa/audit", params={"entity_type": "report"}, headers=_h(TA, "root", "admin"))).json()
    assert audit[0]["metadata"]["format"] == "docx"
