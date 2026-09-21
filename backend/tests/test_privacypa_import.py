"""Import del Registro da Excel/CSV (§67, §68): dry-run, commit transazionale, sicurezza dei file."""
import io
import zipfile
from datetime import date

import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy import func, select

from backend.models.privacypa_assets import Asset
from backend.models.privacypa_catalog import (
    DataCategoryCatalog,
    DataSubjectCategoryCatalog,
    GdprDataClass,
    LegalBasisCatalog,
)
from backend.models.privacypa_identity import PrivacyRoleCode as R, ScopeType
from backend.models.privacypa_import import ImportJob
from backend.models.privacypa_organization import OrganizationalUnit, OrganizationalUnitType
from backend.models.privacypa_parties import Party, PartyType
from backend.models.privacypa_processing import ProcessingActivity
from backend.services.privacypa_audit_service import AuditService
from backend.services.privacypa_import_parser import (
    MAX_ROWS,
    parse_bool,
    parse_date,
    resolve_legal_basis,
    split_list,
    suggest_mapping,
)
from backend.tests.conftest import TA, TB, _h

URL = "/api/privacypa/imports/processing-activities"

HEADERS = ["Titolo", "Struttura", "Descrizione", "Finalità", "Base giuridica", "Interessati",
           "Categorie di dati", "Destinatari", "Responsabili", "Sistemi", "Data inizio",
           "Prossima revisione", "Conservazione", "Responsabile interno"]

GOOD = ["Gestione servizi sociali", "Servizi Sociali", "Interventi sociali", "Istruttoria; Erogazione prestazioni",
        "art. 6 par. 1 lett. e", "Cittadini", "Dati anagrafici; Dati sulla salute", "Fornitore Beta S.r.l.",
        "Fornitore Beta S.r.l.", "Gestionale Servizi Sociali", "01/03/2020", "12/06/2027", "10 anni",
        "Dott.ssa Rossi"]


def make_xlsx(rows, headers=HEADERS, sheet="Registro") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(**over) -> list:
    r = dict(zip(HEADERS, GOOD))
    r.update(over)
    return [r[h] for h in HEADERS]


@pytest_asyncio.fixture
async def imp(env):
    """Anagrafiche di riferimento nel tenant A + un amministratore per l'import."""
    maker = env["maker"]
    async with maker() as db:
        db.add_all([
            OrganizationalUnit(id="u-soc", tenant_id=TA, organization_id=f"org-{TA}", code="SOC",
                               name="Servizi Sociali", unit_type=OrganizationalUnitType.SERVIZIO),
            OrganizationalUnit(id="u-an1", tenant_id=TA, organization_id=f"org-{TA}", code="AN1",
                               name="Anagrafe", unit_type=OrganizationalUnitType.UFFICIO),
            OrganizationalUnit(id="u-an2", tenant_id=TA, organization_id=f"org-{TA}", code="AN2",
                               name="Anagrafe", unit_type=OrganizationalUnitType.UFFICIO),
            OrganizationalUnit(id="u-seg", tenant_id=TB, organization_id=f"org-{TB}", code="SEG",
                               name="Ufficio Segreto", unit_type=OrganizationalUnitType.UFFICIO),
            DataSubjectCategoryCatalog(tenant_id=None, code="CIT", name="Cittadini"),
            DataCategoryCatalog(tenant_id=None, code="ANAG", name="Dati anagrafici", gdpr_class=GdprDataClass.ORDINARY),
            DataCategoryCatalog(tenant_id=None, code="SAL", name="Dati sulla salute", gdpr_class=GdprDataClass.SPECIAL_ART9),
            LegalBasisCatalog(code="GDPR_6_1_E", article="6", paragraph="1", letter="e", description="Compito di interesse pubblico"),
            LegalBasisCatalog(code="GDPR_9_2_G", article="9", paragraph="2", letter="g", description="Interesse pubblico rilevante"),
            LegalBasisCatalog(code="GDPR_6_1_C", article="6", paragraph="1", letter="c", description="Obbligo legale"),
            Party(tenant_id=TA, code="FO-000001", legal_name="Fornitore Beta S.r.l.", party_type=PartyType.PROCESSOR),
            Asset(tenant_id=TA, code="AS-000001", name="Gestionale Servizi Sociali"),
        ])
        await db.commit()
    await env["grant"](TA, "boss", R.PRIVACY_MANAGER)
    return {**env, "who": _h(TA, "boss")}


async def upload(imp, content, name="registro.xlsx", who=None):
    return await imp["client"].post(
        URL, files={"file": (name, content, "application/octet-stream")}, headers=who or imp["who"]
    )


async def dry_run(imp, job_id, mapping, sheet="Registro", who=None, **opts):
    return await imp["client"].post(
        f"{URL}/{job_id}/dry-run", json={"sheet": sheet, "mapping": mapping, **opts}, headers=who or imp["who"]
    )


async def commit(imp, job_id, who=None, confirm=True):
    return await imp["client"].post(f"{URL}/{job_id}/commit", json={"confirm": confirm}, headers=who or imp["who"])


async def flow(imp, rows, **opts):
    """upload + dry-run con mapping suggerito; restituisce (job_id, report)."""
    up = (await upload(imp, make_xlsx(rows))).json()
    mapping = up["sheets"][0]["suggested_mapping"]
    report = (await dry_run(imp, up["job"]["id"], mapping, **opts)).json()
    return up["job"]["id"], report


async def count_processing(imp, tenant=TA) -> int:
    async with imp["maker"]() as db:
        return (await db.execute(
            select(func.count()).select_from(ProcessingActivity).where(ProcessingActivity.tenant_id == tenant)
        )).scalar_one()


# ================================================================= parser (funzioni pure)

@pytest.mark.parametrize("text,expected", [
    ("GDPR_6_1_E", "GDPR_6_1_E"),
    ("gdpr 6 1 e", "GDPR_6_1_E"),
    ("art. 6 par. 1 lett. e", "GDPR_6_1_E"),
    ("Art. 6, paragrafo 1, lettera c)", "GDPR_6_1_C"),
    ("6.1.e", "GDPR_6_1_E"),
    ("9(2)(g)", "GDPR_9_2_G"),
    ("art. 9 lett. g", "GDPR_9_2_G"),
    ("Compito di interesse pubblico", "GDPR_6_1_E"),
    ("art. 6 par. 1 lett. z", None),
    ("consenso libero", None),
    ("art. 99", None),
])
def test_legal_basis_resolution(text, expected):
    catalog = {"GDPR_6_1_E": "Compito di interesse pubblico", "GDPR_9_2_G": "Interesse pubblico rilevante",
               "GDPR_6_1_C": "Obbligo legale", "GDPR_10": "Condanne penali"}
    assert resolve_legal_basis(text, catalog) == expected


def test_value_normalizers():
    assert parse_date("31/12/2025") == date(2025, 12, 31)
    assert parse_date("2025-12-31") == date(2025, 12, 31)
    assert parse_date("31.12.2025") == date(2025, 12, 31)
    assert parse_date("31/02/2025") is None
    assert parse_date("domani") is None
    assert [parse_bool(x) for x in ("Sì", "si", "NO", "x", "boh", None)] == [True, True, False, True, None, None]
    assert split_list("a; b|c\nd, a") == ["a", "b", "c", "d"]
    assert split_list("Consiglio, Giunta", commas=False) == ["Consiglio, Giunta"]
    assert suggest_mapping(HEADERS)["legal_bases"] == "Base giuridica"
    assert suggest_mapping(["Nome trattamento", "Ufficio"]) == {"title": "Nome trattamento", "unit": "Ufficio"}


# ================================================================= flusso

@pytest.mark.asyncio
async def test_upload_returns_sheets_columns_and_suggested_mapping(imp):
    r = await upload(imp, make_xlsx([GOOD]))
    assert r.status_code == 201
    body = r.json()
    sheet = body["sheets"][0]
    assert sheet["name"] == "Registro" and sheet["row_count"] == 1
    assert sheet["columns"] == HEADERS
    assert sheet["suggested_mapping"]["title"] == "Titolo"
    assert sheet["suggested_mapping"]["retention_note"] == "Conservazione"
    assert body["job"]["status"] == "UPLOADED"


@pytest.mark.asyncio
async def test_dry_run_reports_anomalies_and_writes_nothing(imp):
    before = await count_processing(imp)
    rows = [
        GOOD,
        row(Titolo=None),                                                       # titolo mancante → invalida
        row(Titolo="Protocollo", Struttura="Ufficio Inesistente"),              # struttura non trovata
        row(Titolo="Albo", **{"Base giuridica": "consenso libero", "Data inizio": "31/02/2020"}),
        row(Titolo="Tributi", Destinatari="Ente Sconosciuto", **{"Categorie di dati": "Dati esoterici"}),
        row(Titolo="Anagrafe unica", Struttura="Anagrafe"),                     # struttura ambigua
        GOOD,                                                                   # duplicato nel file
    ]
    job_id, report = await flow(imp, rows)
    counts = report["counts"]
    assert counts["rows_total"] == 7
    assert counts["invalid"] == 1 and counts["duplicates"] == 1 and counts["ok"] == 5
    assert counts["can_commit"] is False and counts["blocked_by_invalid_rows"] is True

    by_code = {}
    for a in report["anomalies"]:
        by_code.setdefault(a["code"], []).append(a)
    assert by_code["TITLE_MISSING"][0]["severity"] == "error" and by_code["TITLE_MISSING"][0]["row"] == 3
    assert by_code["UNIT_NOT_FOUND"][0]["value"] == "Ufficio Inesistente"
    assert by_code["UNIT_AMBIGUOUS"][0]["row"] == 7
    assert by_code["LEGAL_BASIS_UNKNOWN"][0]["value"] == "consenso libero"
    assert by_code["DATE_INVALID"][0]["field"] == "start_date"
    assert {a["field"] for a in by_code["REFERENCE_NOT_FOUND"]} == {"recipients", "data_categories"}
    assert by_code["DUPLICATE_IN_FILE"][0]["row"] == 8
    assert report["preview"][0]["legal_bases"] == ["GDPR_6_1_E"]

    assert await count_processing(imp) == before  # nessuna scrittura sul dominio
    async with imp["maker"]() as db:
        actions = {e.action for e in await AuditService(db).list_recent(TA, entity_type="import_job")}
    assert {"upload", "dry_run"} <= actions


@pytest.mark.asyncio
async def test_commit_creates_drafts_with_provenance_and_links(imp):
    parties_before = await imp["maker"]().__aenter__()
    await parties_before.close()
    job_id, report = await flow(imp, [
        GOOD,
        row(Titolo="Sanità", Destinatari="Ente Sconosciuto", **{"Base giuridica": "GDPR_9_2_G; consenso libero"}),
    ])
    assert report["counts"]["can_commit"] is True

    r = await commit(imp, job_id)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "COMMITTED" and body["already_committed"] is False
    assert body["commit"]["created"] == 2

    async with imp["maker"]() as db:
        acts = (await db.execute(
            select(ProcessingActivity).where(ProcessingActivity.tenant_id == TA, ProcessingActivity.source == "LEGACY_IMPORT")
            .order_by(ProcessingActivity.code)
        )).scalars().all()
        n_parties = (await db.execute(select(func.count()).select_from(Party))).scalar_one()
        n_assets = (await db.execute(select(func.count()).select_from(Asset))).scalar_one()
    assert n_parties == 1 and n_assets == 1  # nessuna anagrafica creata dall'import
    assert len(acts) == 2
    first, second = acts
    assert first.status.value == "DRAFT" and first.created_by == "boss"
    assert first.source_reference == f"import:{job_id}:row:2"
    assert first.organizational_unit_id == "u-soc"
    assert first.start_date == date(2020, 3, 1) and first.next_review_date == date(2027, 6, 12)
    assert "Conservazione: 10 anni" in first.legacy_notes and "Responsabile interno: Dott.ssa Rossi" in first.legacy_notes
    assert "Ente Sconosciuto" in second.legacy_notes and "consenso libero" in second.legacy_notes

    detail = (await imp["client"].get(
        f"/api/privacypa/processing-activities/{first.id}", headers=imp["who"])).json()
    assert [b["legal_basis_code"] for b in detail["legal_bases"]] == ["GDPR_6_1_E"]
    assert [p["description"] for p in detail["purposes"]] == ["Istruttoria", "Erogazione prestazioni"]
    assert len(detail["data_subjects"]) == 1 and len(detail["data_categories"]) == 2
    assert len(detail["assets"]) == 1 and len(detail["processors"]) == 1 and len(detail["recipients"]) == 1

    job = (await imp["client"].get(f"{URL}/{job_id}", headers=imp["who"])).json()
    assert job["status"] == "COMMITTED" and job["committed_at"]
    async with imp["maker"]() as db:
        stored = await db.get(ImportJob, job_id)
        assert stored.storage_key is None  # il file caricato viene rimosso a import concluso
        actions = {e.action for e in await AuditService(db).list_recent(TA, entity_type="import_job")}
        created_audit = await AuditService(db).list_for_entity(TA, "processing_activity", first.id)
    assert "commit" in actions and any(e.action == "create" for e in created_audit)


@pytest.mark.asyncio
async def test_commit_requires_dry_run_and_explicit_confirmation(imp):
    up = (await upload(imp, make_xlsx([GOOD]))).json()
    job_id = up["job"]["id"]
    assert (await commit(imp, job_id)).status_code == 409
    await dry_run(imp, job_id, up["sheets"][0]["suggested_mapping"])
    assert (await commit(imp, job_id, confirm=False)).status_code == 422
    assert await count_processing(imp) == 0 + await count_processing(imp)  # invariato dalla conferma mancata
    assert (await commit(imp, job_id)).status_code == 200


@pytest.mark.asyncio
async def test_invalid_rows_block_commit_unless_explicitly_skipped(imp):
    rows = [GOOD, row(Titolo=None), row(Titolo="Protocollo")]
    up = (await upload(imp, make_xlsx(rows))).json()
    job_id, mapping = up["job"]["id"], up["sheets"][0]["suggested_mapping"]

    await dry_run(imp, job_id, mapping)
    blocked = await commit(imp, job_id)
    assert blocked.status_code == 422 and "non valide" in blocked.json()["detail"]

    report = (await dry_run(imp, job_id, mapping, skip_invalid_rows=True)).json()
    assert report["counts"]["can_commit"] is True
    done = (await commit(imp, job_id)).json()
    assert done["commit"]["created"] == 2 and done["commit"]["skipped_invalid"] == 1


@pytest.mark.asyncio
async def test_duplicates_are_skipped_by_default_and_importable_on_request(imp):
    job_id, _ = await flow(imp, [GOOD])
    await commit(imp, job_id)

    again = (await flow(imp, [GOOD, row(Titolo="Nuovo")]))
    report = again[1]
    assert report["counts"]["duplicates"] == 1 and report["counts"]["ok"] == 1
    assert any(a["code"] == "DUPLICATE_EXISTING" for a in report["anomalies"])
    assert (await commit(imp, again[0])).json()["commit"]["created"] == 1

    forced_id, forced = await flow(imp, [GOOD], on_duplicate="import")
    assert forced["counts"]["ok"] == 1 and forced["counts"]["duplicates"] == 0
    assert (await commit(imp, forced_id)).json()["commit"]["created"] == 1


@pytest.mark.asyncio
async def test_commit_is_idempotent(imp):
    job_id, _ = await flow(imp, [GOOD, row(Titolo="Altro")])
    first = await commit(imp, job_id)
    assert first.json()["already_committed"] is False
    total = await count_processing(imp)
    second = await commit(imp, job_id)
    assert second.status_code == 200 and second.json()["already_committed"] is True
    assert await count_processing(imp) == total


@pytest.mark.asyncio
async def test_critical_error_rolls_back_the_whole_import(imp, monkeypatch):
    from backend.services.privacypa_processing_service import ProcessingService

    job_id, _ = await flow(imp, [GOOD, row(Titolo="Due"), row(Titolo="Tre"), row(Titolo="Quattro")])
    before = await count_processing(imp)
    original, calls = ProcessingService.create_activity, {"n": 0}

    async def flaky(self, *args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("guasto simulato")
        return await original(self, *args, **kwargs)

    monkeypatch.setattr(ProcessingService, "create_activity", flaky)
    r = await commit(imp, job_id)
    assert r.status_code == 500
    assert "nessun trattamento è stato creato" in r.json()["detail"] and "guasto" not in r.text

    assert await count_processing(imp) == before  # neanche le prime due righe sono rimaste
    job = (await imp["client"].get(f"{URL}/{job_id}", headers=imp["who"])).json()
    assert job["status"] == "FAILED"
    async with imp["maker"]() as db:
        actions = {e.action for e in await AuditService(db).list_recent(TA, entity_type="import_job")}
    assert "commit_failed" in actions and "commit" not in actions

    # riprovabile: dopo il guasto, dry-run e commit vanno a buon fine
    monkeypatch.setattr(ProcessingService, "create_activity", original)
    mapping = job["mapping"]
    await dry_run(imp, job_id, mapping)
    assert (await commit(imp, job_id)).status_code == 200
    assert await count_processing(imp) == before + 4


@pytest.mark.asyncio
async def test_registry_changes_after_dry_run_require_a_new_dry_run(imp):
    job_id, _ = await flow(imp, [GOOD])
    # nel frattempo qualcuno inserisce lo stesso trattamento: il piano approvato non è più valido
    async with imp["maker"]() as db:
        from backend.services.privacypa_processing_service import ProcessingService
        await ProcessingService(db).create_activity(
            TA, "x", "Gestione servizi sociali", organizational_unit_id="u-soc")
        await db.commit()
    r = await commit(imp, job_id)
    assert r.status_code == 409 and "rieseguire il dry-run" in r.json()["detail"]


# ================================================================= permessi e isolamento

@pytest.mark.asyncio
async def test_import_requires_management_role_with_tenant_wide_scope(imp):
    await imp["grant"](TA, "op", R.OPERATOR)
    await imp["grant"](TA, "scoped", R.PRIVACY_MANAGER, ScopeType.UNIT, "u-soc")
    content = make_xlsx([GOOD])
    for who in (_h(TA, "stranger"), _h(TA, "op"), _h(TA, "scoped")):
        assert (await upload(imp, content, who=who)).status_code == 403
    assert (await upload(imp, content, who=_h(TA, "root", "admin"))).status_code == 201


@pytest.mark.asyncio
async def test_jobs_are_tenant_isolated_and_units_are_not_resolved_across_tenants(imp):
    up = (await upload(imp, make_xlsx([row(Titolo="Segreto", Struttura="Ufficio Segreto")]))).json()
    job_id, mapping = up["job"]["id"], up["sheets"][0]["suggested_mapping"]
    report = (await dry_run(imp, job_id, mapping)).json()
    assert any(a["code"] == "UNIT_NOT_FOUND" for a in report["anomalies"])  # l'unità esiste solo nel tenant B

    other = _h(TB, "root", "admin")
    assert (await imp["client"].get(f"{URL}/{job_id}", headers=other)).status_code == 404
    assert (await imp["client"].get(f"{URL}/{job_id}/sheets", headers=other)).status_code == 404
    assert (await dry_run(imp, job_id, mapping, who=other)).status_code == 404
    assert (await commit(imp, job_id, who=other)).status_code == 404
    assert (await imp["client"].delete(f"{URL}/{job_id}", headers=other)).status_code == 404
    assert (await imp["client"].get(URL, headers=other)).json() == []


@pytest.mark.asyncio
async def test_discard_removes_the_uploaded_file(imp, tmp_path):
    up = (await upload(imp, make_xlsx([GOOD]))).json()
    files = [p for p in (tmp_path / "imports").rglob("*") if p.is_file()]
    assert len(files) == 1 and files[0].name == up["job"]["id"]
    r = await imp["client"].delete(f"{URL}/{up['job']['id']}", headers=imp["who"])
    assert r.json()["status"] == "DISCARDED"
    assert not [p for p in (tmp_path / "imports").rglob("*") if p.is_file()]
    assert (await dry_run(imp, up["job"]["id"], up["sheets"][0]["suggested_mapping"])).status_code == 409


@pytest.mark.asyncio
async def test_committed_job_cannot_be_discarded(imp):
    job_id, _ = await flow(imp, [GOOD])
    await commit(imp, job_id)
    assert (await imp["client"].delete(f"{URL}/{job_id}", headers=imp["who"])).status_code == 409


# ================================================================= mapping e formati

@pytest.mark.asyncio
@pytest.mark.parametrize("mapping,fragment", [
    ({"unit": "Struttura"}, "titolo"),
    ({"title": "Titolo", "campo_inventato": "Struttura"}, "non riconosciuti"),
    ({"title": "Colonna che non c'è"}, "non presenti"),
    ({"title": "Titolo", "unit": "Titolo"}, "più campi"),
])
async def test_mapping_is_validated(imp, mapping, fragment):
    up = (await upload(imp, make_xlsx([GOOD]))).json()
    r = await dry_run(imp, up["job"]["id"], mapping)
    assert r.status_code == 422 and fragment in r.json()["detail"]


@pytest.mark.asyncio
async def test_csv_with_semicolons_cp1252_dates_and_booleans(imp):
    csv_text = (
        "Nome trattamento;Ufficio;Finalità;Base giuridica;Data inizio;Profilazione;Conferimento obbligatorio\r\n"
        "Gestione città;Servizi Sociali;Verifica dell'utenza;6.1.e;15/01/2021;Sì;No\r\n"
    )
    up = (await upload(imp, csv_text.encode("cp1252"), name="registro.csv")).json()
    assert up["sheets"][0]["name"] == "CSV"
    job_id = up["job"]["id"]
    await dry_run(imp, job_id, up["sheets"][0]["suggested_mapping"], sheet="CSV")
    assert (await commit(imp, job_id)).status_code == 200
    async with imp["maker"]() as db:
        act = (await db.execute(select(ProcessingActivity).where(ProcessingActivity.title == "Gestione città"))).scalar_one()
    assert act.profiling is True and act.mandatory_provision is False
    assert act.start_date == date(2021, 1, 15) and act.organizational_unit_id == "u-soc"


@pytest.mark.asyncio
async def test_cells_are_read_as_data_and_formulas_are_never_evaluated(imp):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    ws.append(HEADERS)
    ws.append(GOOD)
    literal = row(Titolo="placeholder")
    ws.append(literal)
    ws.cell(row=3, column=1).value = '=HYPERLINK("http://evil","clic")'
    ws.cell(row=3, column=1).data_type = "s"  # testo che inizia con "=", non una formula
    ws.append(row(Titolo="placeholder"))
    ws.cell(row=4, column=1).value = "=1+1"  # formula vera: nel file non ha un valore calcolato
    buf = io.BytesIO()
    wb.save(buf)

    up = (await upload(imp, buf.getvalue())).json()
    job_id, mapping = up["job"]["id"], up["sheets"][0]["suggested_mapping"]
    report = (await dry_run(imp, job_id, mapping, skip_invalid_rows=True)).json()
    assert report["counts"]["invalid"] == 1  # la formula non viene eseguita: il titolo risulta vuoto
    assert [a["row"] for a in report["anomalies"] if a["code"] == "TITLE_MISSING"] == [4]

    assert (await commit(imp, job_id)).status_code == 200
    async with imp["maker"]() as db:
        titles = (await db.execute(select(ProcessingActivity.title))).scalars().all()
    assert '=HYPERLINK("http://evil","clic")' in titles  # conservato come testo, mai interpretato


def _zip_with(names_and_data: dict[str, bytes], compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as zf:
        for name, data in names_and_data.items():
            zf.writestr(name, data)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_hostile_or_unsupported_files_are_rejected(imp):
    valid = make_xlsx([GOOD])
    macro = _zip_with({"xl/workbook.xml": b"<x/>", "xl/vbaProject.bin": b"MZ"})
    bomb = _zip_with({"xl/workbook.xml": b"<x/>", "xl/big.bin": b"0" * (101 * 1024 * 1024)})
    many_rows = make_xlsx([[f"T{i}"] + [None] * 13 for i in range(MAX_ROWS + 1)])
    cases = [
        ("vecchio.xls", valid),
        ("dati.txt", b"Titolo\nA"),
        ("finto.xlsx", b"non sono uno zip"),
        ("macro.xlsx", macro),
        ("bomba.xlsx", bomb),
        ("vuoto.xlsx", b""),
        ("binario.csv", b"Titolo;Struttura\n\x00\x01\x02;x"),
        ("solo_header.csv", b"Titolo;Struttura\n"),
        ("enorme.xlsx", many_rows),
        ("grande.csv", b"a;b\n" + b"x;y\n" * (3 * 1024 * 1024)),
    ]
    for name, content in cases:
        r = await upload(imp, content, name=name)
        assert r.status_code == 422, (name, r.status_code, r.text[:200])
    async with imp["maker"]() as db:
        assert (await db.execute(select(func.count()).select_from(ImportJob))).scalar_one() == 0


@pytest.mark.asyncio
async def test_client_filename_is_never_used_as_a_path(imp, tmp_path):
    up = (await upload(imp, make_xlsx([GOOD]), name="../../../evil.xlsx")).json()
    assert "/" not in up["job"]["filename"] and "\\" not in up["job"]["filename"]
    stored = [p for p in tmp_path.rglob("*") if p.is_file()]
    assert len(stored) == 1 and stored[0].name == up["job"]["id"]
    assert stored[0].resolve().is_relative_to((tmp_path / "imports").resolve())
