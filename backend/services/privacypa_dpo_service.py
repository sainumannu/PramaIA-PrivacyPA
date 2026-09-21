"""
PrivacyPA — Pareri del DPO (§43), DPO workspace (§43) e relazione annuale (§42).

Il parere è un record distinto dalla decisione del titolare: una volta emesso non si modifica. La relazione
annuale genera dati e statistiche dell'anno; commenti e conclusioni restano modificabili dal soggetto competente
finché non viene finalizzata, e a quel punto anche i dati calcolati vengono congelati.
"""
from collections import Counter
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.auth.privacypa_scope import VisibilityScope
from backend.models.privacypa_assessments import Dpia, DpiaStatus
from backend.models.privacypa_compliance import (
    ActionStatus,
    AuditStatus,
    BreachStatus,
    CorrectiveAction,
    CorrectiveSource,
    DataBreach,
    DataSubjectRequest,
    PrivacyAudit,
    RequestStatus,
)
from backend.models.privacypa_catalog import DataCategoryCatalog, GdprDataClass
from backend.models.privacypa_parties import PartyType, Subprocessor
from backend.models.privacypa_people import (
    OPINION_ENTITY_TYPES,
    AnnualReport,
    AnnualReportStatus,
    AssignmentStatus,
    DpoOpinion,
    OpinionStatus,
    TrainingAssignment,
    TrainingCompletion,
    TrainingCourse,
)
from backend.models.privacypa_processing import ProcessingActivity, ProcessingDataCategory, ProcessingStatus
from backend.models.privacypa_retention import RetentionRule
from backend.models.privacypa_rules import FindingStatus, Severity
from backend.services.privacypa_approval_service import ApprovalService
from backend.services.privacypa_compliance_service import ComplianceError, ComplianceNotFound, _Base
from backend.services.privacypa_findings_service import Context, agreement_is_valid
from backend.services.privacypa_rules_engine import RulesEngine

SECTIONS = (
    "registry", "reviews", "dpia", "requests", "breaches", "audits", "suppliers", "training", "actions",
    "criticalities", "dpo_activity", "next_year_plan",
)
COMMENT_SECTIONS = SECTIONS + ("conclusions",)
MAX_COMMENT = 8000


# ================================================================ pareri

class OpinionService(_Base):
    def _q(self, tenant_id: str):
        return select(DpoOpinion).where(DpoOpinion.tenant_id == tenant_id)

    async def get(self, tenant_id: str, opinion_id: str) -> DpoOpinion | None:
        return (await self.db.execute(self._q(tenant_id).where(DpoOpinion.id == opinion_id))).scalar_one_or_none()

    async def search(
        self, tenant_id: str, *, entity_type: str | None = None, entity_id: str | None = None,
        status: OpinionStatus | None = None,
    ) -> list[DpoOpinion]:
        stmt = self._q(tenant_id)
        if entity_type:
            stmt = stmt.where(DpoOpinion.entity_type == entity_type)
        if entity_id:
            stmt = stmt.where(DpoOpinion.entity_id == entity_id)
        if status:
            stmt = stmt.where(DpoOpinion.status == status)
        return list((await self.db.execute(stmt.order_by(DpoOpinion.requested_at.desc(), DpoOpinion.id))).scalars().all())

    async def entity_ref(self, tenant_id: str, entity_type: str, entity_id: str) -> tuple[str | None, str | None, str | None]:
        """Codice, titolo e struttura dell'elemento (None se non esiste); ComplianceError se il tipo non ammette pareri."""
        if entity_type not in OPINION_ENTITY_TYPES:
            raise ComplianceError(f"Tipo di elemento non ammesso: {entity_type}")
        if entity_type == "data_breach":
            from backend.services.privacypa_compliance_service import BreachService

            b = await BreachService(self.db).get(tenant_id, entity_id)
            return (b.code, b.title, b.organizational_unit_id) if b else (None, None, None)
        info = await ApprovalService(self.db).load_entity(tenant_id, entity_type, entity_id)
        return (info.code, info.title, info.unit_id) if info else (None, None, None)

    async def request(self, tenant_id: str, actor: str, entity_type: str, entity_id: str, notes: str | None,
                      dpo_user_id: str | None = None) -> DpoOpinion:
        code, title, unit = await self.entity_ref(tenant_id, entity_type, entity_id)
        if code is None:
            raise ComplianceError("Elemento non trovato")
        pending = (await self.db.execute(self._q(tenant_id).where(
            DpoOpinion.entity_type == entity_type, DpoOpinion.entity_id == entity_id,
            DpoOpinion.status == OpinionStatus.REQUESTED))).first()
        if pending:
            raise ComplianceError("Esiste già una richiesta di parere in attesa per questo elemento")
        opinion = DpoOpinion(
            tenant_id=tenant_id, entity_type=entity_type, entity_id=entity_id, entity_code=code, entity_title=title,
            organizational_unit_id=unit, requested_by=actor, request_notes=(notes or "").strip() or None,
            dpo_user_id=(dpo_user_id or "").strip() or None,
        )
        self.db.add(opinion)
        await self.db.flush()
        await self.audit.log(tenant_id, "dpo_opinion", "request", event_type="dpo_opinion.requested", entity_id=opinion.id,
                             actor_user_id=actor, new_values={"entity_type": entity_type, "entity_id": entity_id, "status": "REQUESTED"})
        return opinion

    async def issue(self, tenant_id: str, opinion_id: str, actor: str, text: str) -> DpoOpinion:
        opinion = await self.get(tenant_id, opinion_id)
        if opinion is None:
            raise ComplianceNotFound("Parere non trovato")
        if opinion.status != OpinionStatus.REQUESTED:
            raise ComplianceError("Il parere è già stato emesso o ritirato e non si modifica")
        if not (text or "").strip():
            raise ComplianceError("Il testo del parere è obbligatorio")
        opinion.opinion_text, opinion.issued_at, opinion.dpo_user_id = text.strip(), datetime.utcnow(), actor
        opinion.status = OpinionStatus.ISSUED
        await self.db.flush()
        await self.audit.log(tenant_id, "dpo_opinion", "issue", event_type="dpo_opinion.issued", entity_id=opinion.id,
                             actor_user_id=actor, new_values={"entity_type": opinion.entity_type, "entity_id": opinion.entity_id, "status": "ISSUED"})
        return opinion

    async def withdraw(self, tenant_id: str, opinion_id: str, actor: str) -> DpoOpinion:
        opinion = await self.get(tenant_id, opinion_id)
        if opinion is None:
            raise ComplianceNotFound("Parere non trovato")
        if opinion.status != OpinionStatus.REQUESTED:
            raise ComplianceError("Si ritira solo una richiesta ancora in attesa")
        opinion.status = OpinionStatus.WITHDRAWN
        await self.db.flush()
        await self.audit.log(tenant_id, "dpo_opinion", "withdraw", event_type="dpo_opinion.withdrawn", entity_id=opinion.id,
                             actor_user_id=actor, new_values={"status": "WITHDRAWN"})
        return opinion


# ================================================================ workspace

def _short(items: list, limit: int = 25) -> list:
    return items[:limit]


class DpoWorkspaceService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def build(self, tenant_id: str, scope: VisibilityScope, user_id: str, today: date | None = None) -> dict:
        today = today or date.today()
        now = datetime.utcnow()
        engine = RulesEngine(self.db)
        await engine.evaluate(tenant_id, today)

        async def rows(stmt):
            return list((await self.db.execute(stmt)).scalars().all())

        review = [p for p in await rows(select(ProcessingActivity).where(
            ProcessingActivity.tenant_id == tenant_id, ProcessingActivity.deleted_at.is_(None),
            ProcessingActivity.status == ProcessingStatus.IN_REVIEW)) if scope.unit_visible(p.organizational_unit_id)]
        dpias = [d for d in await rows(select(Dpia).where(
            Dpia.tenant_id == tenant_id, Dpia.deleted_at.is_(None),
            Dpia.status.in_((DpiaStatus.DPO_REVIEW, DpiaStatus.AWAITING_DECISION)))) if scope.dpia_visible(d)]
        opinions = [o for o in await OpinionService(self.db).search(tenant_id, status=OpinionStatus.REQUESTED)
                    if scope.unit_visible(o.organizational_unit_id) and o.dpo_user_id in (None, user_id)]
        breaches = [b for b in await rows(select(DataBreach).where(
            DataBreach.tenant_id == tenant_id, DataBreach.deleted_at.is_(None), DataBreach.status != BreachStatus.CLOSED)
            .order_by(DataBreach.aware_at)) if scope.unit_visible(b.organizational_unit_id) and scope.level_ok(b.confidentiality_level, b.created_by)]
        audits = [a for a in await rows(select(PrivacyAudit).where(
            PrivacyAudit.tenant_id == tenant_id, PrivacyAudit.deleted_at.is_(None),
            PrivacyAudit.status.in_((AuditStatus.PLANNED, AuditStatus.IN_PROGRESS))).order_by(PrivacyAudit.planned_date))
            if scope.unit_visible(a.organizational_unit_id)]
        actions = [a for a in await rows(select(CorrectiveAction).where(
            CorrectiveAction.tenant_id == tenant_id, CorrectiveAction.deleted_at.is_(None),
            CorrectiveAction.source_type == CorrectiveSource.DPO_RECOMMENDATION,
            CorrectiveAction.status.in_((ActionStatus.OPEN, ActionStatus.IN_PROGRESS))).order_by(CorrectiveAction.due_date))
            if scope.unit_visible(a.organizational_unit_id)]
        open_findings = [f for f in await engine.list_findings(tenant_id, statuses=(FindingStatus.OPEN, FindingStatus.ACKNOWLEDGED))
                         if (f.processing_activity_id and scope.processing_visible(f.processing_activity_id))
                         or (not f.processing_activity_id and (not f.organizational_unit_id or scope.unit_visible(f.organizational_unit_id)))]
        requests = [r for r in await rows(select(DataSubjectRequest).where(
            DataSubjectRequest.tenant_id == tenant_id, DataSubjectRequest.deleted_at.is_(None),
            DataSubjectRequest.status.in_((RequestStatus.RECEIVED, RequestStatus.IN_PROGRESS))))
            if scope.unit_visible(r.organizational_unit_id)]
        overdue_actions = [a for a in await rows(select(CorrectiveAction).where(
            CorrectiveAction.tenant_id == tenant_id, CorrectiveAction.deleted_at.is_(None),
            CorrectiveAction.status.in_((ActionStatus.OPEN, ActionStatus.IN_PROGRESS)), CorrectiveAction.due_date < today))
            if scope.unit_visible(a.organizational_unit_id)]

        horizon = today + timedelta(days=60)
        calendar = []
        for a in audits:
            if a.planned_date <= horizon:
                calendar.append({"date": a.planned_date.isoformat(), "kind": "audit", "entity_type": "audit_plan", "entity_id": a.id, "code": a.code})
        for p in await rows(select(ProcessingActivity).where(
            ProcessingActivity.tenant_id == tenant_id, ProcessingActivity.deleted_at.is_(None),
            ProcessingActivity.next_review_date.is_not(None), ProcessingActivity.next_review_date <= horizon,
            ProcessingActivity.status.in_((ProcessingStatus.APPROVED, ProcessingStatus.ACTIVE)))):
            if scope.unit_visible(p.organizational_unit_id):
                calendar.append({"date": p.next_review_date.isoformat(), "kind": "processing_review", "entity_type": "processing_activity", "entity_id": p.id, "code": p.code})
        for r in requests:
            if r.due_at.date() <= horizon:
                calendar.append({"date": r.due_at.date().isoformat(), "kind": "request_deadline", "entity_type": "data_subject_request", "entity_id": r.id, "code": r.code})
        for a in actions + overdue_actions:
            if a.due_date <= horizon:
                calendar.append({"date": a.due_date.isoformat(), "kind": "action_due", "entity_type": "corrective_action", "entity_id": a.id, "code": a.code})
        seen, unique = set(), []
        for item in sorted(calendar, key=lambda i: (i["date"], i["kind"], i["code"] or "")):
            key = (item["kind"], item["entity_id"])
            if key not in seen:
                seen.add(key)
                unique.append(item)

        report = (await self.db.execute(select(AnnualReport).where(AnnualReport.tenant_id == tenant_id, AnnualReport.year == today.year))).scalar_one_or_none()
        item = lambda e, **extra: {"id": e.id, "code": getattr(e, "code", None) or getattr(e, "entity_code", None), **extra}  # noqa: E731
        return {
            "generated_at": now.isoformat(),
            "processing_in_review": [item(p, title=p.title, unit_id=p.organizational_unit_id) for p in _short(review)],
            "dpias_to_review": [item(d, title=d.title, status=d.status.value) for d in _short(dpias)],
            "opinions_requested": [
                {"id": o.id, "entity_type": o.entity_type, "entity_id": o.entity_id, "code": o.entity_code, "title": o.entity_title,
                 "requested_at": o.requested_at.isoformat(), "assigned_to_me": o.dpo_user_id == user_id}
                for o in _short(opinions)
            ],
            "breaches_open": [item(b, title=b.title, status=b.status.value, hours_since_aware=round((now - b.aware_at).total_seconds() / 3600, 1),
                                   notify_authority=b.notify_authority) for b in _short(breaches)],
            "audits": [item(a, title=a.title, status=a.status.value, planned_date=a.planned_date.isoformat()) for a in _short(audits)],
            "recommended_actions": [item(a, title=a.title, due_date=a.due_date.isoformat(), status=a.status.value) for a in _short(actions)],
            "followups": {
                "findings_open": len(open_findings),
                "findings_high": sum(1 for f in open_findings if f.severity == Severity.HIGH),
                "requests_open": len(requests),
                "requests_overdue": sum(1 for r in requests if r.due_at < now),
                "actions_overdue": len(overdue_actions),
                "opinions_pending": len(opinions),
            },
            "calendar": _short(unique, 40),
            "annual_report": {"year": today.year, "status": report.status.value if report else None},
        }


# ================================================================ relazione annuale

def _year_range(year: int) -> tuple[datetime, datetime]:
    return datetime(year, 1, 1), datetime(year + 1, 1, 1)


async def compute_annual_report(db: AsyncSession, tenant_id: str, year: int) -> dict:
    """Dati e statistiche dell'anno: nessun testo di commento, solo numeri verificabili."""
    start, end = _year_range(year)
    today = date.today()
    now = datetime.utcnow()

    async def rows(stmt):
        return list((await db.execute(stmt)).scalars().all())

    def in_year(moment) -> bool:
        return moment is not None and start <= (moment if isinstance(moment, datetime) else datetime.combine(moment, datetime.min.time())) < end

    processing = await rows(select(ProcessingActivity).where(ProcessingActivity.tenant_id == tenant_id, ProcessingActivity.deleted_at.is_(None)))
    special = set((await db.execute(
        select(ProcessingDataCategory.processing_activity_id).join(
            DataCategoryCatalog, DataCategoryCatalog.id == ProcessingDataCategory.data_category_id,
        ).where(DataCategoryCatalog.gdpr_class.in_((GdprDataClass.SPECIAL_ART9, GdprDataClass.CRIMINAL_ART10)))
    )).scalars().all())
    with_rules = set((await db.execute(select(RetentionRule.processing_activity_id).where(
        RetentionRule.tenant_id == tenant_id, RetentionRule.deleted_at.is_(None)))).scalars().all())
    live = [p for p in processing if p.status in (ProcessingStatus.APPROVED, ProcessingStatus.ACTIVE)]

    dpias = await rows(select(Dpia).where(Dpia.tenant_id == tenant_id, Dpia.deleted_at.is_(None)))
    requests = await rows(select(DataSubjectRequest).where(DataSubjectRequest.tenant_id == tenant_id, DataSubjectRequest.deleted_at.is_(None)))
    req_year = [r for r in requests if in_year(r.received_at)]
    answered = [r for r in req_year if r.response_sent_at]
    breaches = [b for b in await rows(select(DataBreach).where(DataBreach.tenant_id == tenant_id, DataBreach.deleted_at.is_(None))) if in_year(b.aware_at)]
    audits = await rows(select(PrivacyAudit).where(PrivacyAudit.tenant_id == tenant_id, PrivacyAudit.deleted_at.is_(None)))
    audits_year = [a for a in audits if a.planned_date and in_year(a.planned_date)]
    actions = await rows(select(CorrectiveAction).where(CorrectiveAction.tenant_id == tenant_id, CorrectiveAction.deleted_at.is_(None)))
    open_actions = [a for a in actions if a.status in (ActionStatus.OPEN, ActionStatus.IN_PROGRESS)]

    await RulesEngine(db).evaluate(tenant_id, today)
    ctx = Context(db, tenant_id, today)
    parties = list((await ctx.parties()).values())
    by_party = await ctx.agreements_by_party()
    processors = [p for p in parties if p.party_type == PartyType.PROCESSOR]
    next_start, next_end = _year_range(year + 1)
    expiring_next = [
        a for a in await ctx.agreements()
        if a.expiry_date and next_start.date() <= a.expiry_date < next_end.date()
    ]
    subprocessors = await rows(select(Subprocessor).where(Subprocessor.tenant_id == tenant_id, Subprocessor.deleted_at.is_(None)))

    courses = await rows(select(TrainingCourse).where(TrainingCourse.tenant_id == tenant_id, TrainingCourse.deleted_at.is_(None)))
    mandatory_ids = {c.id for c in courses if c.mandatory}
    assignments = [a for a in await rows(select(TrainingAssignment).where(TrainingAssignment.tenant_id == tenant_id))
                   if a.course_id in mandatory_ids and in_year(a.due_date)]
    completions = [c for c in await rows(select(TrainingCompletion).where(TrainingCompletion.tenant_id == tenant_id)) if in_year(c.completed_at)]
    done = [a for a in assignments if a.status == AssignmentStatus.COMPLETED]
    trainings_expiring = [c for c in await rows(select(TrainingCompletion).where(TrainingCompletion.tenant_id == tenant_id))
                          if c.next_due_at and next_start.date() <= c.next_due_at < next_end.date()]

    opinions = await OpinionService(db).search(tenant_id)
    findings = await RulesEngine(db).list_findings(tenant_id, statuses=(FindingStatus.OPEN, FindingStatus.ACKNOWLEDGED))
    crit = Counter((f.rule_code, f.severity.value) for f in findings)

    def avg(values):
        values = [v for v in values if v is not None]
        return round(sum(values) / len(values), 1) if values else None

    return {
        "year": year,
        "generated_at": now.isoformat(),
        "registry": {
            "total": len(processing), "by_status": dict(Counter(p.status.value for p in processing)),
            "created_in_year": sum(1 for p in processing if in_year(p.created_at)),
            "with_special_categories": sum(1 for p in processing if p.id in special),
            "live_without_retention": sum(1 for p in live if p.id not in with_rules),
        },
        "reviews": {
            "due_in_year": sum(1 for p in live if in_year(p.next_review_date)),
            "reviewed_in_year": sum(1 for p in processing if in_year(p.last_review_date)),
            "overdue_now": sum(1 for p in live if p.next_review_date and p.next_review_date < today),
        },
        "dpia": {
            "total": len(dpias), "by_status": dict(Counter(d.status.value for d in dpias)),
            "created_in_year": sum(1 for d in dpias if in_year(d.created_at)),
            "approved_in_year": sum(1 for d in dpias if in_year(d.approved_at)),
        },
        "requests": {
            "received": len(req_year), "by_type": dict(Counter(r.request_type.value for r in req_year)),
            "responded": len(answered), "extended": sum(1 for r in req_year if r.extension_applied),
            "within_deadline": sum(1 for r in answered if r.response_sent_at <= r.due_at),
            "avg_response_days": avg([(r.response_sent_at - r.received_at).total_seconds() / 86400 for r in answered]),
            "overdue_open": sum(1 for r in requests if r.status in (RequestStatus.RECEIVED, RequestStatus.IN_PROGRESS) and r.due_at < now),
        },
        "breaches": {
            "registered": len(breaches), "by_risk": dict(Counter(b.risk_level.value for b in breaches if b.risk_level)),
            "notified_to_authority": sum(1 for b in breaches if b.notify_authority and b.authority_notified_at),
            "not_notified_with_rationale": sum(1 for b in breaches if b.notify_authority is False),
            "without_decision": sum(1 for b in breaches if b.notify_authority is None),
            "closed": sum(1 for b in breaches if b.status == BreachStatus.CLOSED),
            "avg_hours_to_notification": avg([(b.authority_notified_at - b.aware_at).total_seconds() / 3600 for b in breaches if b.authority_notified_at]),
        },
        "audits": {
            "planned": len(audits_year), "completed": sum(1 for a in audits_year if a.status == AuditStatus.COMPLETED),
            "findings_by_severity": dict(Counter(f.severity for a in audits_year for f in a.findings)),
            "open_findings": sum(1 for a in audits for f in a.findings if f.status.value != "CLOSED"),
        },
        "suppliers": {
            "parties": len(parties), "processors": len(processors),
            "processors_without_valid_agreement": sum(1 for p in processors if not any(agreement_is_valid(a, today) for a in by_party.get(p.id, []))),
            "agreements_expiring_next_year": len(expiring_next), "subprocessors": len(subprocessors),
        },
        "training": {
            "courses": len(courses), "mandatory_assignments": len(assignments), "completed": len(done),
            "overdue": sum(1 for a in assignments if a.status == AssignmentStatus.ASSIGNED and a.due_date < today),
            "completions_in_year": len(completions),
            "completion_rate": round(100 * len(done) / len(assignments), 1) if assignments else None,
        },
        "actions": {
            "opened_in_year": sum(1 for a in actions if in_year(a.created_at)),
            "completed_in_year": sum(1 for a in actions if a.status == ActionStatus.COMPLETED and in_year(a.closed_at)),
            "open": len(open_actions), "overdue": sum(1 for a in open_actions if a.due_date < today),
        },
        "criticalities": [
            {"rule_code": code, "severity": severity, "count": count}
            for (code, severity), count in sorted(crit.items(), key=lambda kv: ({"HIGH": 0, "MEDIUM": 1, "LOW": 2}[kv[0][1]], -kv[1], kv[0][0]))
        ][:15],
        "dpo_activity": {
            "opinions_requested": sum(1 for o in opinions if in_year(o.requested_at)),
            "opinions_issued": sum(1 for o in opinions if in_year(o.issued_at)),
            "opinions_pending": sum(1 for o in opinions if o.status == OpinionStatus.REQUESTED),
        },
        "next_year_plan": {
            "audits_planned": sum(1 for a in audits if a.planned_date and next_start.date() <= a.planned_date < next_end.date() and a.status == AuditStatus.PLANNED),
            "actions_due": sum(1 for a in open_actions if next_start.date() <= a.due_date < next_end.date()),
            "reviews_due": sum(1 for p in live if p.next_review_date and next_start.date() <= p.next_review_date < next_end.date()),
            "agreements_expiring": len(expiring_next), "trainings_expiring": len(trainings_expiring),
        },
    }


class AnnualReportService(_Base):
    async def get(self, tenant_id: str, year: int) -> AnnualReport | None:
        return (await self.db.execute(select(AnnualReport).where(AnnualReport.tenant_id == tenant_id, AnnualReport.year == year))).scalar_one_or_none()

    @staticmethod
    def _check_year(year: int) -> None:
        if not 2000 <= year <= date.today().year + 1:
            raise ComplianceError("Anno non valido")

    async def view(self, tenant_id: str, year: int) -> dict:
        """Relazione dell'anno: se finalizzata i dati congelati, altrimenti quelli calcolati adesso."""
        self._check_year(year)
        report = await self.get(tenant_id, year)
        final = report is not None and report.status == AnnualReportStatus.FINAL and report.snapshot
        return {
            "year": year,
            "status": report.status.value if report else AnnualReportStatus.DRAFT.value,
            "row_version": report.row_version if report else 0,
            "comments": (report.comments if report else None) or {},
            "finalized_by": report.finalized_by if report else None,
            "finalized_at": report.finalized_at.isoformat() if report and report.finalized_at else None,
            "watermark": None if final else "bozza",
            "data": report.snapshot if final else await compute_annual_report(self.db, tenant_id, year),
        }

    async def _locked(self, tenant_id: str, year: int, version: int) -> AnnualReport | None:
        report = await self.get(tenant_id, year)
        current = report.row_version if report else 0
        if current != version:
            raise VersionConflict(current)
        return report

    async def save_comments(self, tenant_id: str, year: int, actor: str, comments: dict, version: int) -> AnnualReport:
        self._check_year(year)
        unknown = set(comments) - set(COMMENT_SECTIONS)
        if unknown:
            raise ComplianceError(f"Sezioni sconosciute: {', '.join(sorted(unknown))}")
        if any(len(v or "") > MAX_COMMENT for v in comments.values()):
            raise ComplianceError(f"Ogni commento è limitato a {MAX_COMMENT} caratteri")
        report = await self._locked(tenant_id, year, version)
        if report is not None and report.status == AnnualReportStatus.FINAL:
            raise ComplianceError("La relazione è finalizzata: riaprirla per modificare i commenti")
        clean = {k: v.strip() for k, v in comments.items() if (v or "").strip()}
        if report is None:
            report = AnnualReport(tenant_id=tenant_id, year=year, comments=clean, updated_by=actor)
            self.db.add(report)
        else:
            report.comments, report.updated_by, report.updated_at = clean, actor, datetime.utcnow()
        await self.db.flush()
        await self.audit.log(tenant_id, "annual_report", "comments", event_type="annual_report.comments.updated", entity_id=report.id,
                             actor_user_id=actor, new_values={"year": year, "sections": sorted(clean)})
        return report

    async def finalize(self, tenant_id: str, year: int, actor: str, version: int) -> AnnualReport:
        self._check_year(year)
        report = await self._locked(tenant_id, year, version)
        if report is not None and report.status == AnnualReportStatus.FINAL:
            raise ComplianceError("La relazione è già finalizzata")
        snapshot = await compute_annual_report(self.db, tenant_id, year)
        if report is None:
            report = AnnualReport(tenant_id=tenant_id, year=year, comments={})
            self.db.add(report)
        report.status, report.snapshot = AnnualReportStatus.FINAL, snapshot
        report.finalized_by, report.finalized_at, report.updated_by, report.updated_at = actor, datetime.utcnow(), actor, datetime.utcnow()
        await self.db.flush()
        await self.audit.log(tenant_id, "annual_report", "finalize", event_type="annual_report.finalized", entity_id=report.id,
                             actor_user_id=actor, new_values={"year": year, "status": "FINAL"})
        return report

    async def reopen(self, tenant_id: str, year: int, actor: str, version: int) -> AnnualReport:
        report = await self._locked(tenant_id, year, version)
        if report is None or report.status != AnnualReportStatus.FINAL:
            raise ComplianceError("La relazione non è finalizzata")
        report.status, report.snapshot, report.finalized_by, report.finalized_at = AnnualReportStatus.DRAFT, None, None, None
        report.updated_by, report.updated_at = actor, datetime.utcnow()
        await self.db.flush()
        await self.audit.log(tenant_id, "annual_report", "reopen", event_type="annual_report.reopened", entity_id=report.id,
                             actor_user_id=actor, new_values={"year": year, "status": "DRAFT"})
        return report


class VersionConflict(Exception):
    def __init__(self, current: int):
        super().__init__("La relazione è stata modificata da un altro utente. Ricaricare e riprovare.")
        self.current = current


# ================================================================ documento

SECTION_TITLES = {
    "registry": "1. Stato del registro", "reviews": "2. Revisioni", "dpia": "3. DPIA", "requests": "4. Richieste degli interessati",
    "breaches": "5. Violazioni dei dati", "audits": "6. Audit", "suppliers": "7. Fornitori", "training": "8. Formazione",
    "actions": "9. Azioni correttive", "criticalities": "10. Principali criticità", "dpo_activity": "11. Attività del DPO",
    "next_year_plan": "12. Piano dell'anno successivo",
}
_LABELS = {
    "total": "Totale", "by_status": "Per stato", "created_in_year": "Creati nell'anno", "with_special_categories": "Con categorie particolari",
    "live_without_retention": "Approvati/attivi senza conservazione", "due_in_year": "Da rivedere nell'anno", "reviewed_in_year": "Rivisti nell'anno",
    "overdue_now": "In ritardo oggi", "approved_in_year": "Approvate nell'anno", "received": "Ricevute", "by_type": "Per tipo",
    "responded": "Evase", "extended": "Con proroga", "within_deadline": "Evase nei termini", "avg_response_days": "Giorni medi di risposta",
    "overdue_open": "Aperte oltre il termine", "registered": "Registrate", "by_risk": "Per rischio", "notified_to_authority": "Notificate all'Autorità",
    "not_notified_with_rationale": "Non notificate con motivazione", "without_decision": "Senza decisione di notifica", "closed": "Chiuse",
    "avg_hours_to_notification": "Ore medie dalla conoscenza alla notifica", "planned": "Pianificati", "completed": "Completati",
    "findings_by_severity": "Rilievi per gravità", "open_findings": "Rilievi aperti", "parties": "Soggetti", "processors": "Responsabili",
    "processors_without_valid_agreement": "Responsabili senza atto valido", "agreements_expiring_next_year": "Atti in scadenza l'anno dopo",
    "subprocessors": "Sub-responsabili", "courses": "Corsi", "mandatory_assignments": "Assegnazioni obbligatorie", "overdue": "In ritardo",
    "completions_in_year": "Completamenti nell'anno", "completion_rate": "Tasso di completamento (%)", "opened_in_year": "Aperte nell'anno",
    "completed_in_year": "Completate nell'anno", "open": "Aperte", "opinions_requested": "Pareri richiesti", "opinions_issued": "Pareri emessi",
    "opinions_pending": "Pareri in attesa", "audits_planned": "Audit pianificati", "actions_due": "Azioni in scadenza",
    "reviews_due": "Revisioni in scadenza", "agreements_expiring": "Atti art. 28 in scadenza", "trainings_expiring": "Formazioni in scadenza",
}


def _plain(value) -> str:
    if isinstance(value, dict):
        return ", ".join(f"{k}: {v}" for k, v in value.items()) or "—"
    return "—" if value is None else str(value)


def annual_report_to_docx(view: dict) -> bytes:
    """Documento Word della relazione: dati dell'anno e, sotto ogni sezione, il commento del soggetto competente."""
    import io

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"Relazione annuale sulla protezione dei dati - {view['year']}", level=0)
    if view.get("watermark"):
        run = doc.add_paragraph().add_run(f"{view['watermark'].upper()} - relazione non finalizzata")
        run.bold = True
    else:
        doc.add_paragraph(f"Finalizzata il {view.get('finalized_at', '')[:10]}.")
    comments, data = view.get("comments") or {}, view.get("data") or {}
    for key in SECTIONS:
        doc.add_heading(SECTION_TITLES[key], level=1)
        section = data.get(key)
        if isinstance(section, list):
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, label in zip(table.rows[0].cells, ("Regola", "Gravità", "Segnalazioni")):
                cell.text = label
            for item in section:
                cells = table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text = item["rule_code"], item["severity"], str(item["count"])
        elif isinstance(section, dict):
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for name, value in section.items():
                cells = table.add_row().cells
                cells[0].text, cells[1].text = _LABELS.get(name, name), _plain(value)
        if comments.get(key):
            doc.add_paragraph()
            doc.add_paragraph(comments[key])
    if comments.get("conclusions"):
        doc.add_heading("Conclusioni", level=1)
        doc.add_paragraph(comments["conclusions"])
    doc.styles["Normal"].font.size = Pt(10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
