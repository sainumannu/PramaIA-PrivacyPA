"""
Report Registro art. 30: JSON/CSV/PDF (§41, §153).

Il report risolve gli identificativi in nomi leggibili e distingue titolare, DPO,
finalità, interessati, categorie di dati, destinatari, trasferimenti, misure. Watermark
"bozza" se non tutti i trattamenti sono approvati. Include le regole di conservazione (§21); un
trattamento senza regole lo dichiara ("non definita") invece di ometterle in silenzio.
"""
import csv
import io
import textwrap
from datetime import datetime

from sqlalchemy import select

from backend.models.privacypa_assessments import InternationalTransfer, TransferLink
from backend.models.privacypa_assets import Asset
from backend.models.privacypa_compliance import LegalReference, ProcessingLegalReference
from backend.models.privacypa_catalog import (
    DataCategoryCatalog,
    DataSubjectCategoryCatalog,
    LegalBasisCatalog,
)
from backend.models.privacypa_organization import Organization, OrganizationalUnit
from backend.models.privacypa_parties import Party
from backend.models.privacypa_processing import ProcessingStatus
from backend.models.privacypa_risk import MeasureImplementation, SecurityMeasure
from backend.services.privacypa_asset_service import AssetService
from backend.services.privacypa_party_service import PartyService
from backend.services.privacypa_processing_service import ProcessingService
from backend.services.privacypa_retention_service import DISPOSAL_LABEL, RetentionService, describe_period

_PROCESSING_TYPES = ("processing_activity", "processing")
_SPECIAL_CLASSES = ("SPECIAL_ART9", "CRIMINAL_ART10")


def _value(x):
    return x.value if hasattr(x, "value") else x


async def _lookups(db, tenant_id: str) -> dict:
    """Carica in blocco le anagrafiche necessarie a risolvere gli identificativi."""
    async def rows(stmt):
        return list((await db.execute(stmt)).scalars().all())

    org = (await db.execute(select(Organization).where(Organization.tenant_id == tenant_id))).scalar_one_or_none()
    catalog_scope = lambda model: (model.tenant_id.is_(None)) | (model.tenant_id == tenant_id)  # noqa: E731
    return {
        "org": org,
        "units": {u.id: u for u in await rows(select(OrganizationalUnit).where(OrganizationalUnit.tenant_id == tenant_id))},
        "legal": {b.code: b for b in await rows(select(LegalBasisCatalog))},
        "subjects": {s.id: s for s in await rows(select(DataSubjectCategoryCatalog).where(catalog_scope(DataSubjectCategoryCatalog)))},
        "categories": {c.id: c for c in await rows(select(DataCategoryCatalog).where(catalog_scope(DataCategoryCatalog)))},
        "parties": {p.id: p for p in await rows(select(Party).where(Party.tenant_id == tenant_id))},
        "assets": {a.id: a for a in await rows(select(Asset).where(Asset.tenant_id == tenant_id))},
        "measures": {m.id: m for m in await rows(select(SecurityMeasure).where(
            (SecurityMeasure.tenant_id.is_(None)) | (SecurityMeasure.tenant_id == tenant_id)))},
        "implementations": await rows(select(MeasureImplementation).where(
            MeasureImplementation.tenant_id == tenant_id,
            MeasureImplementation.entity_type.in_(_PROCESSING_TYPES),
        )),
        "transfers": {t.id: t for t in await rows(select(InternationalTransfer).where(
            InternationalTransfer.tenant_id == tenant_id, InternationalTransfer.deleted_at.is_(None)))},
        "transfer_links": await rows(select(TransferLink).where(TransferLink.entity_type.in_(_PROCESSING_TYPES))),
        "retention": await RetentionService(db).list_for_tenant(tenant_id),
        "legal_refs": {l.id: l for l in await rows(select(LegalReference).where(LegalReference.tenant_id == tenant_id))},
        "processing_legal_refs": await rows(select(ProcessingLegalReference)),
    }


def _resolve(record: dict, activity_id: str, lk: dict) -> dict:
    unit = lk["units"].get(record.get("organizational_unit_id"))
    legal = [
        {"code": b["legal_basis_code"], "description": getattr(lk["legal"].get(b["legal_basis_code"]), "description", None)}
        for b in record.get("legal_bases") or []
    ]
    subjects = [
        getattr(lk["subjects"].get(s["category_id"]), "name", None) or s["category_id"]
        for s in record.get("data_subjects") or []
    ]
    categories = []
    for c in record.get("data_categories") or []:
        cat = lk["categories"].get(c["data_category_id"])
        categories.append({
            "name": getattr(cat, "name", None) or c["data_category_id"],
            "gdpr_class": _value(getattr(cat, "gdpr_class", None)),
        })
    recipients = [
        {
            "name": getattr(lk["parties"].get(r["party_id"]), "legal_name", None) or r["party_id"],
            "communication_type": r.get("communication_type"),
        }
        for r in record.get("recipients") or []
    ]
    processors = [
        {"name": getattr(lk["parties"].get(p["party_id"]), "legal_name", None) or p["party_id"]}
        for p in record.get("processors") or []
    ]
    systems = [
        {"name": getattr(lk["assets"].get(a["asset_id"]), "name", None) or a["asset_id"], "role": a.get("role")}
        for a in record.get("assets") or []
    ]
    transfer_ids = {l.transfer_id for l in lk["transfer_links"] if l.entity_id == activity_id}
    transfers = [
        {
            "code": t.code, "title": t.title, "country_code": t.country_code,
            "mechanism": _value(t.mechanism), "status": _value(t.status),
            "importer": getattr(lk["parties"].get(t.importer_party_id), "legal_name", None),
        }
        for tid in sorted(transfer_ids) if (t := lk["transfers"].get(tid))
    ]
    measures = [
        {
            "name": getattr(lk["measures"].get(i.measure_id), "name", None) or i.measure_id,
            "status": _value(i.implementation_status),
        }
        for i in lk["implementations"] if i.entity_id == activity_id
    ]
    retention = [
        {
            "category": getattr(lk["categories"].get(r.data_category_id), "name", None) if r.data_category_id else None,
            "record_class": r.record_class,
            "period": describe_period(r.retention_unit, r.retention_value, r.start_trigger),
            "criterion": r.retention_criterion,
            "legal_reference": getattr(lk["legal_refs"].get(r.legal_reference_id), "citation", None) or r.legal_reference,
            "disposal": DISPOSAL_LABEL[r.disposal_method],
        }
        for r in lk["retention"] if r.processing_activity_id == activity_id
    ]
    return {
        "unit_name": unit.name if unit else None,
        "retention_resolved": retention,
        "legal_bases_resolved": legal,
        "legal_references_resolved": [
            lk["legal_refs"][l.legal_reference_id].citation
            for l in lk["processing_legal_refs"] if l.processing_activity_id == activity_id and l.legal_reference_id in lk["legal_refs"]
        ],
        "data_subject_names": subjects,
        "data_categories_resolved": categories,
        "has_special_categories": any(c["gdpr_class"] in _SPECIAL_CLASSES for c in categories),
        "recipients_resolved": recipients,
        "processors_resolved": processors,
        "systems": systems,
        "international_transfers": transfers,
        "security_measures": measures,
    }


async def collect_art30(db, tenant_id: str, accessible_unit_ids: set[str] | None, extracted_by: str) -> dict:
    proc_svc = ProcessingService(db)
    asset_svc = AssetService(db)
    party_svc = PartyService(db)
    lk = await _lookups(db, tenant_id)
    activities = await proc_svc.list_activities(tenant_id, accessible_unit_ids=accessible_unit_ids)
    records = []
    for a in activities:
        full = await proc_svc.get_activity(tenant_id, a.id)
        if not full:
            continue
        record = {
            **full.to_dict(include_relations=True),
            "assets": [x.to_dict() for x in await asset_svc.list_links(a.id)],
            "processors": [x.to_dict() for x in await party_svc.list_processors(a.id)],
            "recipients": [x.to_dict() for x in await party_svc.list_recipients(a.id)],
        }
        record.update(_resolve(record, a.id, lk))
        records.append(record)
    approved = all(
        r.get("status") in (ProcessingStatus.APPROVED.value, ProcessingStatus.ACTIVE.value, ProcessingStatus.ARCHIVED.value)
        for r in records
    ) if records else False
    org = lk["org"]
    return {
        "report": "registro_art_30",
        "schema_version": "2.0",
        "extracted_at": datetime.utcnow().isoformat(),
        "extracted_by": extracted_by,
        "tenant_id": tenant_id,
        "controller": {
            "name": org.name, "tax_code": org.tax_code, "ipa_code": org.ipa_code, "pec": org.pec,
            "dpo_name": org.dpo_name, "dpo_contact": org.dpo_contact,
        } if org else None,
        "filters": {"scope_limited": accessible_unit_ids is not None},
        "unavailable_sections": [],
        "watermark": None if approved else "bozza",
        "count": len(records),
        "processing_activities": records,
    }


def _csv_safe(value) -> str:
    """Neutralizza le formule (CSV injection) nei valori che iniziano con = + - @."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _join(items) -> str:
    return " | ".join(str(i) for i in items if i)


def _legal_text(record: dict) -> str:
    """Basi giuridiche GDPR seguite dai riferimenti normativi specifici (§13)."""
    bases = _join(f"{b['code']} {b.get('description') or ''}".strip() for b in record.get("legal_bases_resolved") or [])
    refs = _join(record.get("legal_references_resolved") or [])
    return f"{bases} | Riferimenti: {refs}" if refs and bases else (bases or refs)


def _retention_text(rule: dict) -> str:
    scope = rule.get("category") or rule.get("record_class")
    parts = [f"{scope}: " if scope else "", f"{rule['period']} ({rule['criterion']})", f" - {rule['disposal']}"]
    if rule.get("legal_reference"):
        parts.append(f" [{rule['legal_reference']}]")
    return "".join(parts)


def art30_to_csv(payload: dict) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    controller = payload.get("controller") or {}
    writer.writerow([
        "codice", "titolo", "stato", "versione", "struttura", "finalita",
        "basi_giuridiche", "interessati", "categorie_dati", "categorie_particolari",
        "destinatari", "responsabili", "sistemi", "trasferimenti_extra_ue", "misure",
        "conservazione", "data_inizio", "prossima_revisione",
    ])
    for r in payload.get("processing_activities", []):
        purposes = _join(p.get("description") for p in r.get("purposes") or []) or (r.get("purpose_summary") or "")
        row = [
            r.get("code"), r.get("title"), r.get("status"), r.get("version_no"),
            r.get("unit_name") or "",
            purposes,
            _join(f"{b['code']} {b.get('description') or ''}".strip() for b in r.get("legal_bases_resolved") or []),
            _join(r.get("data_subject_names") or []),
            _join(c["name"] for c in r.get("data_categories_resolved") or []),
            "sì" if r.get("has_special_categories") else "no",
            _join(f"{x['name']} ({x.get('communication_type')})" for x in r.get("recipients_resolved") or []),
            _join(x["name"] for x in r.get("processors_resolved") or []),
            _join(x["name"] for x in r.get("systems") or []),
            _join(f"{t['code']} {t['country_code']} {t.get('mechanism')}" for t in r.get("international_transfers") or []),
            _join(f"{m['name']} [{m['status']}]" for m in r.get("security_measures") or []),
            _join(_retention_text(x) for x in r.get("retention_resolved") or []) or "non definita",
            r.get("start_date") or "", r.get("next_review_date") or "",
        ]
        writer.writerow([_csv_safe(v) for v in row])
    writer.writerow([])
    writer.writerow(["titolare", _csv_safe(controller.get("name")), "dpo", _csv_safe(controller.get("dpo_name")),
                     "contatto_dpo", _csv_safe(controller.get("dpo_contact"))])
    writer.writerow(["estratto_il", payload.get("extracted_at"), "da", _csv_safe(payload.get("extracted_by")),
                     "schema", payload.get("schema_version"), "watermark", payload.get("watermark") or ""])
    if payload.get("unavailable_sections"):
        writer.writerow(["sezioni_non_disponibili", _join(payload["unavailable_sections"])])
    return buf.getvalue()


def art30_to_xlsx(payload: dict) -> bytes:
    """Cartella Excel: `Registro` (una riga per trattamento), `Conservazione` (una per regola), `Informazioni`."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def put(ws, row: list) -> None:
        # Valori sempre come testo: openpyxl tratta come formula ogni stringa che inizia con "=".
        ws.append([_csv_safe(v) if isinstance(v, str) else v for v in row])

    def finish(ws, widths: list[int]) -> None:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    put(ws, ["Codice", "Titolo", "Stato", "Versione", "Struttura", "Finalità", "Basi giuridiche", "Interessati",
             "Categorie di dati", "Categorie particolari", "Destinatari", "Responsabili", "Sistemi",
             "Trasferimenti extra UE", "Misure", "Conservazione", "Data inizio", "Prossima revisione"])
    records = payload.get("processing_activities", [])
    for r in records:
        purposes = _join(p.get("description") for p in r.get("purposes") or []) or (r.get("purpose_summary") or "")
        put(ws, [
            r.get("code"), r.get("title"), r.get("status"), r.get("version_no"), r.get("unit_name") or "", purposes,
            _legal_text(r),
            _join(r.get("data_subject_names") or []),
            _join(c["name"] for c in r.get("data_categories_resolved") or []),
            "sì" if r.get("has_special_categories") else "no",
            _join(f"{x['name']} ({x.get('communication_type')})" for x in r.get("recipients_resolved") or []),
            _join(x["name"] for x in r.get("processors_resolved") or []),
            _join(x["name"] for x in r.get("systems") or []),
            _join(f"{t['code']} {t['country_code']} {t.get('mechanism')}" for t in r.get("international_transfers") or []),
            _join(f"{m['name']} [{m['status']}]" for m in r.get("security_measures") or []),
            _join(_retention_text(x) for x in r.get("retention_resolved") or []) or "non definita",
            r.get("start_date") or "", r.get("next_review_date") or "",
        ])
    finish(ws, [12, 36, 12, 9, 24, 40, 30, 28, 32, 12, 32, 28, 28, 28, 36, 48, 12, 14])

    ret = wb.create_sheet("Conservazione")
    put(ret, ["Trattamento", "Categoria / classe", "Termine", "Criterio", "Riferimento normativo", "Metodo a fine ciclo"])
    for r in records:
        for x in r.get("retention_resolved") or []:
            put(ret, [r.get("code"), x.get("category") or x.get("record_class") or "intero trattamento",
                      x["period"], x["criterion"], x.get("legal_reference") or "", x["disposal"]])
    finish(ret, [12, 30, 34, 48, 28, 26])

    controller = payload.get("controller") or {}
    info = wb.create_sheet("Informazioni")
    put(info, ["Voce", "Valore"])
    for label, value in (
        ("Titolare", controller.get("name")), ("Codice fiscale", controller.get("tax_code")), ("PEC", controller.get("pec")),
        ("DPO", controller.get("dpo_name")), ("Contatto DPO", controller.get("dpo_contact")),
        ("Estratto il", payload.get("extracted_at")), ("Estratto da", payload.get("extracted_by")),
        ("Versione schema", payload.get("schema_version")), ("Trattamenti", payload.get("count")),
        ("Stato documento", (payload.get("watermark") or "").upper() + " - non tutti i trattamenti sono approvati" if payload.get("watermark") else "approvato"),
        ("Estratto limitato alle strutture dell'utente", "sì" if payload.get("filters", {}).get("scope_limited") else "no"),
    ):
        put(info, [label, value if value is not None else ""])
    finish(info, [42, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def art30_to_docx(payload: dict) -> bytes:
    """Documento Word: intestazione con titolare e DPO, poi una scheda (tabella a due colonne) per trattamento."""
    from docx import Document
    from docx.shared import Pt

    controller = payload.get("controller") or {}
    doc = Document()
    doc.add_heading("Registro dei trattamenti (art. 30 GDPR)", level=0)
    if payload.get("watermark"):
        p = doc.add_paragraph()
        run = p.add_run(f"{payload['watermark'].upper()} - documento non approvato")
        run.bold = True
    head = doc.add_table(rows=0, cols=2)
    head.style = "Table Grid"
    for label, value in (
        ("Titolare", controller.get("name")), ("PEC", controller.get("pec")), ("DPO", controller.get("dpo_name")),
        ("Contatto DPO", controller.get("dpo_contact")), ("Estratto il", payload.get("extracted_at")),
        ("Estratto da", payload.get("extracted_by")), ("Trattamenti", payload.get("count")),
    ):
        cells = head.add_row().cells
        cells[0].text, cells[1].text = label, "" if value is None else str(value)
    if payload.get("filters", {}).get("scope_limited"):
        doc.add_paragraph("Estratto limitato alle strutture dell'utente.")

    for r in payload.get("processing_activities", []):
        doc.add_heading(f"{r.get('code')} - {r.get('title')}", level=1)
        purposes = _join(p.get("description") for p in r.get("purposes") or []) or (r.get("purpose_summary") or "")
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        rows = [
            ("Stato", f"{r.get('status')} (v{r.get('version_no')})"), ("Struttura", r.get("unit_name") or ""),
            ("Finalità", purposes),
            ("Basi giuridiche", _legal_text(r)),
            ("Interessati", _join(r.get("data_subject_names") or [])),
            ("Categorie di dati", _join(
                c["name"] + (" (particolare)" if c.get("gdpr_class") in _SPECIAL_CLASSES else "")
                for c in r.get("data_categories_resolved") or [])),
            ("Destinatari", _join(f"{x['name']} ({x.get('communication_type')})" for x in r.get("recipients_resolved") or [])),
            ("Responsabili", _join(x["name"] for x in r.get("processors_resolved") or [])),
            ("Sistemi", _join(x["name"] for x in r.get("systems") or [])),
            ("Trasferimenti extra UE", _join(f"{t['code']} {t['country_code']} ({t.get('mechanism')})" for t in r.get("international_transfers") or [])),
            ("Misure di sicurezza", _join(f"{m['name']} [{m['status']}]" for m in r.get("security_measures") or [])),
            ("Conservazione", _join(_retention_text(x) for x in r.get("retention_resolved") or []) or "non definita"),
            ("Prossima revisione", r.get("next_review_date") or ""),
        ]
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
        doc.add_paragraph()
    doc.styles["Normal"].font.size = Pt(10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def art30_to_pdf(payload: dict) -> bytes:
    """PDF minimale Type1 Helvetica, testo latin-1 (italiano); le righe lunghe vanno a capo."""
    controller = payload.get("controller") or {}
    lines = [
        "PrivacyPA - Registro dei trattamenti (art. 30 GDPR)",
        f"Titolare: {controller.get('name') or '-'}   DPO: {controller.get('dpo_name') or '-'} {controller.get('dpo_contact') or ''}".rstrip(),
        f"Estratto: {payload.get('extracted_at')}  Utente: {payload.get('extracted_by')}",
        f"Schema: {payload.get('schema_version')}  Record: {payload.get('count')}",
    ]
    if payload.get("filters", {}).get("scope_limited"):
        lines.append("Estratto limitato alle strutture dell'utente")
    if payload.get("watermark"):
        lines.append(f"WATERMARK: {payload['watermark'].upper()} - documento non approvato")
    if payload.get("unavailable_sections"):
        lines.append(f"Sezioni non disponibili: {_join(payload['unavailable_sections'])}")
    lines.append("-" * 80)

    def field(label: str, value: str) -> list[str]:
        if not value:
            return []
        return textwrap.wrap(f"{label}: {value}", width=100, initial_indent="  ", subsequent_indent="      ")

    for r in payload.get("processing_activities", []):
        lines.append(f"{r.get('code')}  {r.get('title')}  [{r.get('status')}]  v{r.get('version_no')}")
        purposes = _join(p.get("description") for p in r.get("purposes") or []) or (r.get("purpose_summary") or "")
        lines += field("Struttura", r.get("unit_name") or "")
        lines += field("Finalita", purposes)
        lines += field("Basi giuridiche", _join(b["code"] for b in r.get("legal_bases_resolved") or []))
        lines += field("Interessati", _join(r.get("data_subject_names") or []))
        lines += field("Categorie di dati", _join(
            c["name"] + (" (particolare)" if c.get("gdpr_class") in _SPECIAL_CLASSES else "")
            for c in r.get("data_categories_resolved") or []))
        lines += field("Destinatari", _join(x["name"] for x in r.get("recipients_resolved") or []))
        lines += field("Responsabili", _join(x["name"] for x in r.get("processors_resolved") or []))
        lines += field("Sistemi", _join(x["name"] for x in r.get("systems") or []))
        lines += field("Trasferimenti extra UE", _join(
            f"{t['country_code']} ({t.get('mechanism')})" for t in r.get("international_transfers") or []))
        lines += field("Misure", _join(m["name"] for m in r.get("security_measures") or []))
        lines += field("Conservazione", _join(_retention_text(x) for x in r.get("retention_resolved") or []) or "non definita")
        lines.append("")
    return _simple_pdf(lines)


def _pdf_escape(text: str) -> str:
    encoded = text.encode("latin-1", "replace").decode("latin-1")
    return encoded.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _simple_pdf(lines: list[str]) -> bytes:
    pages: list[list[str]] = []
    chunk: list[str] = []
    for line in lines:
        chunk.append(line[:110])
        if len(chunk) >= 48:
            pages.append(chunk)
            chunk = []
    if chunk:
        pages.append(chunk)
    if not pages:
        pages = [["(vuoto)"]]

    objects = ["%PDF-1.4"]
    body_ids = []
    # we'll assemble after we know content streams
    streams = []
    for page_lines in pages:
        y = 800
        cmds = ["BT", "/F1 10 Tf"]
        for line in page_lines:
            cmds.append(f"1 0 0 1 40 {y} Tm ({_pdf_escape(line)}) Tj")
            y -= 14
        cmds.append("ET")
        content = "\n".join(cmds).encode("latin-1", "replace")
        streams.append(content)

    # Object graph: 1 catalog, 2 pages, 3 font, then page N, content N
    # Simpler: 1=catalog 2=pages 3=font then pairs
    n_pages = len(streams)
    font_id = 3
    page_ids = list(range(4, 4 + n_pages))
    content_ids = list(range(4 + n_pages, 4 + 2 * n_pages))

    def obj(i: int, payload: bytes) -> bytes:
        return f"{i} 0 obj\n".encode() + payload + b"\nendobj\n"

    kids = " ".join(f"{i} 0 R" for i in page_ids)
    catalog = b"<< /Type /Catalog /Pages 2 0 R >>"
    pages_obj = f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    parts = [b"%PDF-1.4\n"]
    offsets = [0]

    def add(i: int, payload: bytes):
        offsets.append(len(b"".join(parts)))
        parts.append(obj(i, payload))

    add(1, catalog)
    add(2, pages_obj)
    add(3, font)
    for pid, cid, stream in zip(page_ids, content_ids, streams):
        page = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {cid} 0 R >>"
        ).encode()
        add(pid, page)
        stream_obj = f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        add(cid, stream_obj)

    xref_pos = len(b"".join(parts))
    xref = [f"xref\n0 {len(offsets)}\n", "0000000000 65535 f \n"]
    for off in offsets[1:]:
        xref.append(f"{off:010d} 00000 n \n")
    trailer = f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    return b"".join(parts) + "".join(xref).encode() + trailer.encode()
