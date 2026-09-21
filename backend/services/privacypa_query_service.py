"""
PrivacyPA — Query builder (§40).

Un'interrogazione è un tipo di elemento, un modo di combinare le condizioni (tutte / almeno una) e un elenco di
condizioni `campo operatore valore`. Campi e operatori vengono da un catalogo dichiarato qui: nessun nome di colonna o
frammento SQL arriva dall'utente, i valori sono parametri.

Oltre alle colonne del modello ci sono campi calcolati: «flag» (es. il trattamento ha una DPIA: sì/no) e campi a valori
multipli (tag, categorie di dati: la condizione vale se *esiste* un valore che la soddisfa; «diverso da» e «non in» valgono
se non ne esiste nessuno). `not_equals` e `not_in` includono i valori vuoti, come si aspetta chi scrive «stato diverso da ATTIVO».

I risultati passano dallo stesso filtro della ricerca globale (permesso di modulo, struttura, riservatezza) prima di
essere contati o restituiti; l'esportazione richiede un permesso dedicato ed esclude i livelli con blocco dell'export
massivo. Non ci sono campi personali (titolari, richiedenti, riferimenti a persone) e le richieste degli interessati non
si interrogano.
"""
import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import Any, Callable

from sqlalchemy import Boolean, Date, DateTime, Integer, String, and_, exists, func, literal_column, or_, select
from sqlalchemy import Enum as SAEnum

from backend.models.privacypa_annotations import EntityTag, Tag
from backend.models.privacypa_assessments import Dpia, DpiaStatus
from backend.models.privacypa_assets import Asset, ProcessingAsset
from backend.models.privacypa_catalog import DataCategoryCatalog, GdprDataClass
from backend.models.privacypa_compliance import CorrectiveAction, DataBreach, PrivacyAudit
from backend.models.privacypa_notices import NoticeProcessingLink, PrivacyNotice
from backend.models.privacypa_parties import Contract, Party, ProcessingProcessor, ProcessorProfile
from backend.models.privacypa_processing import ProcessingActivity, ProcessingDataCategory
from backend.models.privacypa_projects import Project
from backend.services.privacypa_annotation_service import refs_for_rows
from backend.services.privacypa_report_service import _csv_safe
from backend.services.privacypa_search_service import escape_like

OPERATORS = ("equals", "not_equals", "contains", "in", "not_in", "date_before", "date_after", "is_null", "is_not_null")
_NO_VALUE = ("is_null", "is_not_null")
_LIST = ("in", "not_in")
_BASIC = ("equals", "not_equals", "in", "not_in", "is_null", "is_not_null")
TYPE_OPERATORS: dict[str, tuple[str, ...]] = {
    "text": ("equals", "not_equals", "contains", "in", "not_in", "is_null", "is_not_null"),
    "id": _BASIC,
    "enum": _BASIC,
    "number": _BASIC,
    "bool": ("equals", "not_equals", "is_null", "is_not_null"),
    "date": ("equals", "date_before", "date_after", "is_null", "is_not_null"),
    "datetime": ("equals", "date_before", "date_after", "is_null", "is_not_null"),
    "flag": ("equals", "not_equals"),
}
MAX_CONDITIONS = 20
MAX_LIST_VALUES = 50
MAX_TEXT = 255
MAX_PAGE = 200
MAX_EXPORT_ROWS = 10000
# Filtrabili ma non mostrati come colonna: testi lunghi che possono contenere dati personali.
_LONG_TEXT = {"description", "summary"}


class QueryError(ValueError):
    pass


@dataclass(frozen=True)
class QField:
    type: str                                              # text|id|enum|number|bool|date|datetime|flag
    expr: Any = None                                       # colonna o sotto-query scalare
    values: tuple[str, ...] | None = None
    enum_class: Any = None
    flag: Callable[[], Any] | None = None                  # espressione EXISTS correlata (campo «flag»)
    multi: Callable[[], tuple[Any, list]] | None = None    # (colonna dei valori, condizioni di collegamento)
    show: bool = True

    @property
    def operators(self) -> tuple[str, ...]:
        return TYPE_OPERATORS[self.type]

    @property
    def sortable(self) -> bool:
        return self.expr is not None and self.multi is None


def _infer(name: str, column) -> QField:
    kind = column.type
    if isinstance(kind, SAEnum):
        return QField("enum", column, tuple(e.value for e in kind.enum_class), kind.enum_class)
    show = name not in _LONG_TEXT
    if isinstance(kind, Boolean):
        return QField("bool", column)
    if isinstance(kind, DateTime):
        return QField("datetime", column)
    if isinstance(kind, Date):
        return QField("date", column)
    if isinstance(kind, Integer):
        return QField("number", column)
    if isinstance(kind, String):
        return QField("id" if name.endswith("_id") else "text", column, show=show)
    raise TypeError(f"Tipo di colonna non supportato per {name}")


@dataclass(frozen=True)
class QEntity:
    model: type
    module: str
    fields: dict[str, QField]
    columns: tuple[str, ...]                       # colonne predefinite dei risultati
    title: str = "title"
    base: tuple = ()                               # condizioni sempre applicate

    def base_where(self, tenant_id: str) -> list:
        clauses = [self.model.tenant_id == tenant_id, *self.base]
        if hasattr(self.model, "deleted_at"):
            clauses.append(self.model.deleted_at.is_(None))
        return clauses


def _tag_field(entity_type: str, model) -> QField:
    return QField("text", multi=lambda: (Tag.name, [
        EntityTag.tag_id == Tag.id, EntityTag.entity_type == entity_type, EntityTag.entity_id == model.id,
        EntityTag.tenant_id == model.tenant_id]))


def _entity(entity_type: str, model, module: str, names: str, columns: str, extra: dict | None = None, **kw) -> QEntity:
    fields = {n: _infer(n, getattr(model, n)) for n in names.split()}
    fields.update(extra or {})
    fields["tag"] = _tag_field(entity_type, model)
    return QEntity(model, module, fields, tuple(columns.split()), **kw)


def _has(model, *clauses) -> Callable[[], Any]:
    return lambda: exists(select(literal_column("1")).select_from(model).where(*clauses))


_PA = ProcessingActivity
ENTITIES: dict[str, QEntity] = {
    "processing_activity": _entity(
        "processing_activity", _PA, "processing",
        "code title description status organizational_unit_id start_date last_review_date next_review_date mandatory_provision "
        "automated_decision profiling confidentiality_level",
        "code title status organizational_unit_id next_review_date",
        extra={
            "has_dpia": QField("flag", flag=_has(Dpia, Dpia.processing_activity_id == _PA.id, Dpia.deleted_at.is_(None), Dpia.status != DpiaStatus.ARCHIVED)),
            "dpia_status": QField("enum", values=tuple(e.value for e in DpiaStatus), enum_class=DpiaStatus, multi=lambda: (
                Dpia.status, [Dpia.processing_activity_id == _PA.id, Dpia.deleted_at.is_(None)])),
            "has_notice": QField("flag", flag=lambda: exists(select(literal_column("1")).select_from(NoticeProcessingLink).join(
                PrivacyNotice, PrivacyNotice.id == NoticeProcessingLink.privacy_notice_id).where(
                NoticeProcessingLink.processing_activity_id == _PA.id, PrivacyNotice.deleted_at.is_(None)))),
            "has_processor": QField("flag", flag=_has(ProcessingProcessor, ProcessingProcessor.processing_activity_id == _PA.id)),
            "has_system": QField("flag", flag=_has(ProcessingAsset, ProcessingAsset.processing_activity_id == _PA.id)),
            "data_category": QField("text", multi=lambda: (DataCategoryCatalog.code, [
                ProcessingDataCategory.data_category_id == DataCategoryCatalog.id, ProcessingDataCategory.processing_activity_id == _PA.id])),
            "data_class": QField("enum", values=tuple(e.value for e in GdprDataClass), enum_class=GdprDataClass, multi=lambda: (
                DataCategoryCatalog.gdpr_class, [ProcessingDataCategory.data_category_id == DataCategoryCatalog.id,
                                                 ProcessingDataCategory.processing_activity_id == _PA.id])),
        },
        base=(_PA.is_current.is_(True),),
    ),
    "dpia": _entity(
        "dpia", Dpia, "dpia",
        "code title description status screening_outcome processing_activity_id next_review_date approved_at confidentiality_level",
        "code title status processing_activity_id next_review_date",
    ),
    "asset": _entity(
        "asset", Asset, "asset",
        "code name description asset_type hosting_type cloud_service hosting_country mfa_supported logging_supported "
        "encryption_at_rest encryption_in_transit backup_enabled lifecycle_status go_live_date retirement_date owner_unit_id supplier_id",
        "code name asset_type lifecycle_status hosting_type",
        extra={"has_processing": QField("flag", flag=_has(ProcessingAsset, ProcessingAsset.asset_id == Asset.id))},
        title="name",
    ),
    "party": _entity(
        "party", Party, "party",
        "code legal_name party_type country_code",
        "code legal_name party_type country_code",
        extra={"processor_status": QField("enum", select(ProcessorProfile.status).where(ProcessorProfile.party_id == Party.id).scalar_subquery(),
                                          tuple(e.value for e in ProcessorProfile.status.type.enum_class), ProcessorProfile.status.type.enum_class)},
        title="legal_name",
    ),
    "contract": _entity(
        "contract", Contract, "party",
        "code title status start_date end_date renewal_date organizational_unit_id party_id",
        "code title status end_date",
    ),
    "privacy_notice": _entity(
        "privacy_notice", PrivacyNotice, "notice",
        "code title audience status owner_unit_id last_reviewed_at",
        "code title status owner_unit_id last_reviewed_at",
        extra={"has_processing": QField("flag", flag=_has(NoticeProcessingLink, NoticeProcessingLink.privacy_notice_id == PrivacyNotice.id))},
    ),
    "data_breach": _entity(
        "data_breach", DataBreach, "breach",
        "code title description status organizational_unit_id detected_at aware_at risk_level confidentiality_affected integrity_affected "
        "availability_affected approx_data_subjects approx_records notify_authority authority_notified_at notify_subjects "
        "subjects_notified_at closed_at confidentiality_level",
        "code title status aware_at risk_level",
    ),
    "project": _entity(
        "project", Project, "project",
        "code title description status proposing_unit_id supplier_id go_live_date post_review_due_date closed_at uses_cloud uses_ai "
        "profiles_people automated_decisions special_data criminal_data minors systematic_monitoring large_scale "
        "privacy_review_required dpia_screening_required dpia_id confidentiality_level",
        "code title status proposing_unit_id go_live_date",
    ),
    "corrective_action": _entity(
        "corrective_action", CorrectiveAction, "action",
        "code title description source_type priority status due_date organizational_unit_id closed_at",
        "code title status priority due_date",
    ),
    "audit": _entity(
        "audit", PrivacyAudit, "auditplan",
        "code title audit_type status planned_date performed_date organizational_unit_id",
        "code title audit_type status planned_date",
    ),
}
# ---------------------------------------------------------------- valori e condizioni

def _one(name: str, f: QField, raw: Any) -> Any:
    kind = f.type
    if kind in ("text", "id"):
        if not isinstance(raw, str) or not raw.strip():
            raise QueryError(f"Il valore di «{name}» deve essere un testo non vuoto")
        if len(raw.strip()) > MAX_TEXT:
            raise QueryError(f"Il valore di «{name}» supera {MAX_TEXT} caratteri")
        return raw.strip()
    if kind == "enum":
        text = str(raw).strip().upper()
        if text not in (f.values or ()):
            raise QueryError(f"Valore non ammesso per «{name}»: {raw!r} (ammessi: {', '.join(f.values or ())})")
        return f.enum_class(text)
    if kind in ("bool", "flag"):
        if isinstance(raw, bool):
            return raw
        if str(raw).strip().lower() in ("true", "si", "sì", "1"):
            return True
        if str(raw).strip().lower() in ("false", "no", "0"):
            return False
        raise QueryError(f"Il valore di «{name}» deve essere sì/no")
    if kind == "number":
        if isinstance(raw, bool):
            raise QueryError(f"Il valore di «{name}» deve essere un numero intero")
        try:
            return int(raw)
        except (TypeError, ValueError):
            raise QueryError(f"Il valore di «{name}» deve essere un numero intero") from None
    if kind in ("date", "datetime"):
        try:
            return date.fromisoformat(str(raw).strip()[:10])
        except ValueError:
            raise QueryError(f"Il valore di «{name}» deve essere una data (AAAA-MM-GG)") from None
    raise QueryError(f"Campo non interrogabile: {name}")


def _date_clause(column, f: QField, op: str, day: date):
    if f.type == "date":
        return {"equals": column == day, "date_before": column < day, "date_after": column > day}[op]
    start = datetime.combine(day, time.min)
    end = start + timedelta(days=1)
    return {"equals": and_(column >= start, column < end), "date_before": column < start, "date_after": column >= end}[op]


def _scalar(column, f: QField, op: str, value: Any):
    if op == "is_null":
        return column.is_(None)
    if op == "is_not_null":
        return column.is_not(None)
    if op in ("date_before", "date_after") or (f.type in ("date", "datetime") and op == "equals"):
        return _date_clause(column, f, op, value)
    text = f.type == "text"
    side = func.lower(column) if text else column
    low = (lambda v: v.lower()) if text else (lambda v: v)
    if op == "contains":
        return column.ilike(f"%{escape_like(value)}%", escape="\\")
    if op == "equals":
        return side == low(value)
    if op == "not_equals":
        return or_(column.is_(None), side != low(value))
    if op == "in":
        return side.in_([low(v) for v in value])
    return or_(column.is_(None), ~side.in_([low(v) for v in value]))          # not_in


def _clause(f: QField, op: str, value: Any):
    if f.type == "flag":
        present = f.flag()
        return present if (value is True) == (op == "equals") else ~present
    if f.multi is not None:
        column, links = f.multi()
        if op in _NO_VALUE:
            found = exists(select(column).where(*links))
            return found if op == "is_not_null" else ~found
        positive = {"not_equals": "equals", "not_in": "in"}.get(op, op)
        found = exists(select(column).where(*links, _scalar(column, f, positive, value)))
        return ~found if op in ("not_equals", "not_in") else found
    return _scalar(f.expr, f, op, value)


def compile_conditions(entity: QEntity, match: str, conditions: list[dict]):
    """Espressione SQL delle condizioni, o `None` se non ce ne sono. QueryError se campo, operatore o valore non sono validi."""
    if match not in ("all", "any"):
        raise QueryError("Il modo di combinazione deve essere «all» o «any»")
    if len(conditions) > MAX_CONDITIONS:
        raise QueryError(f"Al massimo {MAX_CONDITIONS} condizioni")
    parts = []
    for cond in conditions:
        name, op = cond.get("field"), cond.get("operator")
        f = entity.fields.get(name)
        if f is None:
            raise QueryError(f"Campo sconosciuto: {name}")
        if op not in f.operators:
            raise QueryError(f"Operatore «{op}» non ammesso per «{name}» (ammessi: {', '.join(f.operators)})")
        value = None
        if op in _LIST:
            raw = cond.get("value")
            if not isinstance(raw, list) or not raw or len(raw) > MAX_LIST_VALUES:
                raise QueryError(f"«{name}»: serve un elenco di 1-{MAX_LIST_VALUES} valori")
            value = [_one(name, f, v) for v in raw]
        elif op not in _NO_VALUE:
            if cond.get("value") is None:
                raise QueryError(f"«{name}»: manca il valore")
            value = _one(name, f, cond["value"])
        parts.append(_clause(f, op, value))
    if not parts:
        return None
    return and_(*parts) if match == "all" else or_(*parts)


# ---------------------------------------------------------------- esecuzione

def _cell(row, key: str):
    value = getattr(row, key)
    if hasattr(value, "value"):
        return value.value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def describe(entity_key: str) -> dict:
    entity = ENTITIES[entity_key]
    return {
        "key": entity_key, "module": entity.module, "columns": list(entity.columns),
        "fields": [
            {"key": k, "type": f.type, "operators": list(f.operators), "values": list(f.values) if f.values else None,
             "multi": f.multi is not None, "selectable": f.expr is not None and f.multi is None and f.show and f.type != "flag",
             "sortable": f.sortable}
            for k, f in entity.fields.items()
        ],
    }


class QueryService:
    def __init__(self, db):
        self.db = db

    def _entity(self, key: str) -> QEntity:
        if key not in ENTITIES:
            raise QueryError(f"Tipo di elemento non interrogabile: {key}")
        return ENTITIES[key]

    def _columns(self, entity: QEntity, requested: list[str] | None) -> list[str]:
        names = requested or list(entity.columns)
        for n in names:
            f = entity.fields.get(n)
            if f is None or f.expr is None or f.multi is not None or not f.show:
                raise QueryError(f"Colonna non disponibile: {n}")
        return list(dict.fromkeys(names))

    async def _visible(self, tenant_id: str, scope, entity_key: str, match: str, conditions: list[dict], sort: str | None, direction: str) -> list[tuple]:
        entity = self._entity(entity_key)
        clause = compile_conditions(entity, match, conditions)
        model = entity.model
        stmt = select(model).where(*entity.base_where(tenant_id))
        if clause is not None:
            stmt = stmt.where(clause)
        key = sort or "code"
        f = entity.fields.get(key)
        if f is None or not f.sortable:
            raise QueryError(f"Ordinamento non disponibile per: {key}")
        if direction not in ("asc", "desc"):
            raise QueryError("La direzione deve essere «asc» o «desc»")
        order = f.expr.desc() if direction == "desc" else f.expr.asc()
        rows = list((await self.db.execute(stmt.order_by(order, model.id))).scalars().all())
        refs = {r.entity_id: r for r in await refs_for_rows(self.db, entity_key, rows)}
        return [(row, refs[row.id]) for row in rows if refs[row.id].visible_to(scope)]

    async def run(
        self, tenant_id: str, scope, entity_key: str, match: str, conditions: list[dict], *,
        columns: list[str] | None = None, sort: str | None = None, direction: str = "asc", offset: int = 0, limit: int = 50,
    ) -> dict:
        entity = self._entity(entity_key)
        cols = self._columns(entity, columns)
        visible = await self._visible(tenant_id, scope, entity_key, match, conditions, sort, direction)
        page = visible[offset:offset + limit]
        tags = await self._tags(tenant_id, entity_key, [r.id for r, _ in page])
        return {
            "entity": entity_key, "total": len(visible), "offset": offset, "limit": limit,
            "columns": [{"key": c, "type": entity.fields[c].type} for c in cols],
            "rows": [
                {"id": r.id, "code": getattr(r, "code", None), "title": getattr(r, entity.title, None), "tags": tags.get(r.id, []),
                 "values": {c: _cell(r, c) for c in cols}}
                for r, _ in page
            ],
        }

    async def _tags(self, tenant_id: str, entity_key: str, ids: list[str]) -> dict[str, list[str]]:
        if not ids:
            return {}
        found = await self.db.execute(
            select(EntityTag.entity_id, Tag.name).join(Tag, Tag.id == EntityTag.tag_id).where(
                EntityTag.tenant_id == tenant_id, EntityTag.entity_type == entity_key, EntityTag.entity_id.in_(ids)).order_by(Tag.name))
        out: dict[str, list[str]] = {}
        for entity_id, name in found.all():
            out.setdefault(entity_id, []).append(name)
        return out

    async def export_rows(
        self, tenant_id: str, scope, entity_key: str, match: str, conditions: list[dict], *,
        columns: list[str] | None = None, sort: str | None = None, direction: str = "asc",
    ) -> dict:
        """Tutte le righe visibili, meno quelle di un livello di riservatezza con blocco dell'export massivo."""
        entity = self._entity(entity_key)
        cols = self._columns(entity, columns)
        visible = await self._visible(tenant_id, scope, entity_key, match, conditions, sort, direction)
        allowed, excluded = [], 0
        for row, ref in visible:
            rule = scope.levels.rule(ref.level) if (scope.levels is not None and ref.level is not None) else None
            if rule and rule.get("block_bulk_export"):
                excluded += 1
            else:
                allowed.append(row)
        if len(allowed) > MAX_EXPORT_ROWS:
            raise QueryError(f"Troppi risultati da esportare ({len(allowed)}): restringi la ricerca (massimo {MAX_EXPORT_ROWS})")
        tags = await self._tags(tenant_id, entity_key, [r.id for r in allowed])
        return {
            "columns": cols, "excluded": excluded,
            "rows": [[_cell(r, c) for c in cols] + [", ".join(tags.get(r.id, []))] for r in allowed],
        }


def to_csv(columns: list[str], rows: list[list]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([*columns, "tag"])
    for row in rows:
        writer.writerow([_csv_safe(v) if isinstance(v, str) else ("" if v is None else v) for v in row])
    return buf.getvalue()


def to_xlsx(columns: list[str], rows: list[list], info: list[tuple[str, str]]) -> bytes:
    """Foglio `Risultati` e foglio `Informazioni` (ente, data, utente e filtro applicato, §41.3)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Risultati"
    ws.append([*columns, "tag"])
    for row in rows:
        ws.append([_csv_safe(v) if isinstance(v, str) else v for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    meta = wb.create_sheet("Informazioni")
    for label, value in info:
        meta.append([label, _csv_safe(value)])
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 80
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
