"""
PrivacyPA — Lettura e normalizzazione dei file di import del Registro (§67, §68).

Funzioni pure, senza accesso al DB: leggono XLSX/CSV con limiti di sicurezza, propongono il
mapping delle colonne e normalizzano date, booleani, elenchi e basi giuridiche.
"""
import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 60
MAX_CELL_CHARS = 5000
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED = 100 * 1024 * 1024

ALLOWED_EXTENSIONS = (".xlsx", ".csv")

# Campo di destinazione -> intestazioni riconosciute (già normalizzate con `norm`).
FIELD_ALIASES: dict[str, list[str]] = {
    "title": ["titolo", "trattamento", "nome trattamento", "denominazione", "denominazione trattamento",
              "attivita di trattamento", "nome"],
    "unit": ["struttura", "ufficio", "servizio", "settore", "area", "unita organizzativa",
             "struttura organizzativa", "direzione"],
    "description": ["descrizione", "descrizione trattamento", "descrizione del trattamento"],
    "purposes": ["finalita", "finalita del trattamento", "finalita trattamento", "scopo"],
    "legal_bases": ["base giuridica", "basi giuridiche", "liceita", "base di liceita", "condizione di liceita"],
    "data_subjects": ["interessati", "categorie di interessati", "categorie interessati", "soggetti interessati"],
    "data_categories": ["categorie di dati", "categorie dati", "dati trattati", "tipologia di dati",
                        "tipologia dati", "dati personali"],
    "recipients": ["destinatari", "categorie di destinatari", "comunicazione"],
    "processors": ["responsabili", "responsabile esterno", "responsabili del trattamento", "fornitori",
                   "responsabili esterni"],
    "assets": ["sistemi", "applicativi", "software", "applicativo", "sistema", "strumenti"],
    "source_of_data": ["origine dei dati", "fonte dei dati", "provenienza dati"],
    "collection_method": ["modalita di raccolta", "modalita raccolta"],
    "start_date": ["data inizio", "inizio trattamento", "data avvio"],
    "last_review_date": ["ultima revisione", "data ultima revisione", "ultimo aggiornamento"],
    "next_review_date": ["prossima revisione", "data prossima revisione", "revisione entro", "scadenza revisione"],
    "mandatory_provision": ["conferimento obbligatorio", "obbligatorieta conferimento"],
    "automated_decision": ["decisioni automatizzate", "processo decisionale automatizzato"],
    "profiling": ["profilazione"],
    # Dati non modellati: finiscono in `legacy_notes` per non perderli.
    "owner_note": ["responsabile interno", "dirigente", "responsabile", "titolare del trattamento", "referente"],
    "retention_note": ["conservazione", "termini di conservazione", "periodo di conservazione",
                       "tempi di conservazione"],
    "notes": ["note", "annotazioni", "osservazioni"],
}
REQUIRED_FIELDS = ("title",)


class ImportFileError(ValueError):
    """File non leggibile o non ammesso; il messaggio è mostrabile all'utente."""


def norm(text: Any) -> str:
    """Minuscolo, senza accenti né punteggiatura: base per confronti tolleranti."""
    raw = unicodedata.normalize("NFKD", str(text if text is not None else ""))
    ascii_text = raw.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", ascii_text).strip()


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (numero riga nel foglio, valori)

    def as_dicts(self) -> list[tuple[int, dict[str, Any]]]:
        return [(no, dict(zip(self.headers, values))) for no, values in self.rows]


# ----------------------------------------------------------------- lettura file

def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(value)).strip()
    if len(text) > MAX_CELL_CHARS:
        raise ImportFileError(f"Una cella supera {MAX_CELL_CHARS} caratteri")
    return text or None


def _build_sheet(name: str, raw_rows) -> Sheet:
    headers: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []
    for row_no, raw in enumerate(raw_rows, start=1):
        cells = [_clean_cell(c) for c in list(raw)[:MAX_COLUMNS + 1]]
        if len(cells) > MAX_COLUMNS:
            raise ImportFileError(f"Il foglio '{name}' supera {MAX_COLUMNS} colonne")
        if all(c is None for c in cells):
            continue
        if headers is None:
            headers = []
            for i, cell in enumerate(cells, start=1):
                label = str(cell) if cell is not None else f"Colonna {i}"
                base, n = label, 2
                while label in headers:
                    label, n = f"{base} ({n})", n + 1
                headers.append(label)
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(f"Il foglio '{name}' supera {MAX_ROWS} righe: suddividere il file")
        rows.append((row_no, cells + [None] * (len(headers) - len(cells))))
    return Sheet(name=name, headers=headers or [], rows=rows)


def _check_zip(content: bytes) -> zipfile.ZipFile:
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise ImportFileError("Il file non è un XLSX valido")
    infos = zf.infolist()
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED:
        raise ImportFileError("Il file XLSX è troppo grande una volta decompresso")
    names = {i.filename for i in infos}
    if "xl/workbook.xml" not in names or "xl/vbaProject.bin" in names:
        raise ImportFileError("Il file non è un XLSX valido (o contiene macro, non ammesse)")
    return zf


def _read_xlsx(content: bytes) -> list[Sheet]:
    from openpyxl import load_workbook

    _check_zip(content).close()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Impossibile leggere il file XLSX")
    try:
        sheets = [
            _build_sheet(ws.title, ws.iter_rows(values_only=True))
            for ws in wb.worksheets
            if ws.sheet_state == "visible"
        ]
    finally:
        wb.close()
    return [s for s in sheets if s.headers]


def _read_csv(content: bytes) -> list[Sheet]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("Codifica del CSV non riconosciuta (usare UTF-8)")
    if "\x00" in text:
        raise ImportFileError("Il file non sembra un CSV di testo")
    first_line = text.split("\n", 1)[0]
    delimiter = max(";,\t|", key=first_line.count) if any(d in first_line for d in ";,\t|") else ";"
    csv.field_size_limit(MAX_CELL_CHARS + 1)
    try:
        sheet = _build_sheet("CSV", csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error:
        raise ImportFileError("CSV malformato")
    return [sheet] if sheet.headers else []


def read_tabular(content: bytes, filename: str) -> list[Sheet]:
    """Legge un file .xlsx o .csv e restituisce i fogli con intestazioni e righe non vuote."""
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportFileError(f"File troppo grande (max {MAX_IMPORT_BYTES // (1024 * 1024)} MB)")
    if not content:
        raise ImportFileError("File vuoto")
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        sheets = _read_xlsx(content)
    elif lower.endswith(".csv"):
        sheets = _read_csv(content)
    else:
        raise ImportFileError("Formato non supportato: usare .xlsx o .csv (i vecchi .xls vanno salvati come .xlsx)")
    if not sheets:
        raise ImportFileError("Il file non contiene dati")
    if not any(s.rows for s in sheets):
        raise ImportFileError("Il file contiene solo le intestazioni: nessuna riga di dati da importare")
    return sheets


# ----------------------------------------------------------------- mapping colonne

def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Proposta automatica campo -> intestazione, in base agli alias riconosciuti."""
    by_norm = {norm(h): h for h in headers}
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            header = by_norm.get(alias)
            if header and header not in used:
                mapping[field_name] = header
                used.add(header)
                break
    return mapping


def validate_mapping(mapping: dict[str, str], headers: list[str]) -> None:
    unknown = [f for f in mapping if f not in FIELD_ALIASES]
    if unknown:
        raise ImportFileError(f"Campi non riconosciuti nel mapping: {', '.join(unknown)}")
    missing_cols = [c for c in mapping.values() if c not in headers]
    if missing_cols:
        raise ImportFileError(f"Colonne non presenti nel foglio: {', '.join(missing_cols)}")
    if len(set(mapping.values())) != len(mapping):
        raise ImportFileError("La stessa colonna è associata a più campi")
    for required in REQUIRED_FIELDS:
        if required not in mapping:
            raise ImportFileError("La colonna del titolo del trattamento deve essere associata")


# ----------------------------------------------------------------- normalizzatori

_TRUE = {"si", "s", "yes", "y", "true", "vero", "1", "x"}
_FALSE = {"no", "n", "false", "falso", "0"}


def parse_bool(value: Any) -> bool | None:
    token = norm(value)
    if token in _TRUE:
        return True
    if token in _FALSE:
        return False
    return None


_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$"), (0, 1, 2)),
    (re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$"), (2, 1, 0)),
)


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern, (yi, mi, di) in _DATE_PATTERNS:
        m = pattern.match(text)
        if m:
            try:
                return date(int(m.group(yi + 1)), int(m.group(mi + 1)), int(m.group(di + 1)))
            except ValueError:
                return None
    return None


def split_list(value: Any, *, commas: bool = True) -> list[str]:
    pattern = r"[;\n|,]+" if commas else r"[;\n|]+"
    seen: set[str] = set()
    out: list[str] = []
    for part in re.split(pattern, str(value or "")):
        part = part.strip(" \t\r-•*")
        if part and norm(part) not in seen:
            seen.add(norm(part))
            out.append(part)
    return out


_LEGAL_NOISE = re.compile(r"\b(art|articolo|par|paragrafo|comma|lett|lettera|gdpr|rgpd|reg|regolamento|ue|2016|679)\b")


def resolve_legal_basis(token: str, catalog: dict[str, str]) -> str | None:
    """
    Codice del catalogo (es. GDPR_6_1_E) da testo libero. `catalog` mappa codice -> descrizione.
    Riconosce il codice, forme come "art. 6 par. 1 lett. e", "6.1.e", "9(2)(g)", "art. 10"
    e la descrizione esatta del catalogo.
    """
    code = re.sub(r"\s+", "_", token.strip().upper())
    if code in catalog:
        return code
    text = _LEGAL_NOISE.sub(" ", norm(token))
    parts = re.findall(r"\d+|\b[a-j]\b", text)
    candidate = None
    if parts == ["10"]:
        candidate = "GDPR_10"
    elif len(parts) == 3 and parts[0] in ("6", "9") and parts[1] in ("1", "2") and parts[2].isalpha():
        candidate = f"GDPR_{parts[0]}_{parts[1]}_{parts[2].upper()}"
    elif len(parts) == 2 and parts[0] in ("6", "9") and parts[1].isalpha():
        candidate = f"GDPR_{parts[0]}_{'1' if parts[0] == '6' else '2'}_{parts[1].upper()}"
    if candidate in catalog:
        return candidate
    wanted = norm(token)
    for cat_code, description in catalog.items():
        if wanted and wanted == norm(description):
            return cat_code
    return None


def short(value: Any, limit: int = 80) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= limit else text[: limit - 1] + "…"


@dataclass
class Anomaly:
    row: int
    field: str
    code: str
    severity: str  # "error" | "warning"
    message: str
    value: str | None = None

    def to_dict(self) -> dict:
        return {"row": self.row, "field": self.field, "code": self.code, "severity": self.severity,
                "message": self.message, "value": self.value}


@dataclass
class NormalizedRow:
    row_no: int
    data: dict = field(default_factory=dict)
    anomalies: list[Anomaly] = field(default_factory=list)
    outcome: str = "ok"  # ok | invalid | duplicate
