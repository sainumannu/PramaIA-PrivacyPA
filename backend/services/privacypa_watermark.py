"""
PrivacyPA — Filigrana utente/data sui documenti riservati (§135).

Solo PDF e Word (.docx): sono i formati che si possono marcare senza alterarne il contenuto. Il file originale non si
modifica: la copia scaricata porta la filigrana e l'audit conserva lo sha256 dell'originale. Altri formati (immagini,
Excel, archivi...) non si marcano: con la modalità `required` non si scaricano, con `if_possible` si scaricano intatti e
l'audit registra che la filigrana non è stata applicata.
"""
import io
import math
from datetime import datetime

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class WatermarkUnsupported(Exception):
    """Formato non marcabile (o file cifrato/illeggibile)."""


def supported_kind(mime_type: str, file_name: str) -> str | None:
    name = (file_name or "").lower()
    if mime_type == PDF or name.endswith(".pdf"):
        return "pdf"
    if mime_type == DOCX or name.endswith(".docx"):
        return "docx"
    return None


def watermark_text(level: str, user: str, code: str, when: datetime | None = None) -> str:
    """Testo ASCII (il font di base del PDF non ha tutti i caratteri): livello, elemento, utente e momento del download."""
    when = when or datetime.utcnow()
    return f"{level} - {code} - {user} - {when.strftime('%Y-%m-%d %H:%M')} UTC"


def _pdf_escape(text: str) -> bytes:
    safe = text.encode("latin-1", "replace").decode("latin-1")
    return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1")


def _overlay(width: float, height: float, text: str):
    """Pagina trasparente con il testo in diagonale (tenue) e una riga a piè di pagina."""
    from pypdf import PageObject
    from pypdf.generic import DecodedStreamObject, DictionaryObject, FloatObject, NameObject

    page = PageObject.create_blank_page(width=width, height=height)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    faded = DictionaryObject({NameObject("/Type"): NameObject("/ExtGState"), NameObject("/ca"): FloatObject(0.18), NameObject("/CA"): FloatObject(0.18)})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/WM"): font}),
        NameObject("/ExtGState"): DictionaryObject({NameObject("/GS1"): faded}),
    })
    diagonal = math.hypot(width, height)
    size = max(8.0, min(40.0, 0.8 * diagonal / (max(len(text), 1) * 0.55)))
    angle = math.atan2(height, width)
    cos, sin = math.cos(angle), math.sin(angle)
    text_width = len(text) * size * 0.5
    x, y = width / 2 - (text_width / 2) * cos, height / 2 - (text_width / 2) * sin
    escaped = _pdf_escape(text)
    content = (
        b"q /GS1 gs 0.4 0.4 0.4 rg BT /WM %.2f Tf %.4f %.4f %.4f %.4f %.2f %.2f Tm (" % (size, cos, sin, -sin, cos, x, y)
        + escaped + b") Tj ET Q\n"
        + b"q 0.3 0.3 0.3 rg BT /WM 7 Tf 1 0 0 1 20 12 Tm (" + escaped + b") Tj ET Q\n"
    )
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = stream
    return page


def _watermark_pdf(data: bytes, text: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise WatermarkUnsupported("PDF cifrato")
        writer = PdfWriter()
        for page in reader.pages:
            page.merge_page(_overlay(float(page.mediabox.width), float(page.mediabox.height), text))
            writer.add_page(page)
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue()
    except WatermarkUnsupported:
        raise
    except Exception as exc:                       # file danneggiato o con strutture che non si sanno unire
        raise WatermarkUnsupported(f"PDF non marcabile: {exc}") from exc


def _watermark_docx(data: bytes, text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    try:
        doc = Document(io.BytesIO(data))
        for section in doc.sections:
            for part in (section.header, section.footer):
                part.is_linked_to_previous = False
                run = part.add_paragraph().add_run(text)
                run.bold, run.font.size, run.font.color.rgb = True, Pt(8), RGBColor(0xB0, 0x30, 0x30)
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception as exc:
        raise WatermarkUnsupported(f"Documento Word non marcabile: {exc}") from exc


def apply_watermark(data: bytes, mime_type: str, file_name: str, text: str) -> bytes:
    """Copia marcata del file; WatermarkUnsupported se il formato non si può marcare."""
    kind = supported_kind(mime_type, file_name)
    if kind == "pdf":
        return _watermark_pdf(data, text)
    if kind == "docx":
        return _watermark_docx(data, text)
    raise WatermarkUnsupported("Formato non marcabile")
